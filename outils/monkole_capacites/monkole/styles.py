"""Charte graphique reprise de la version retenue (couleurs, polices, formats)."""
from functools import lru_cache

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Synthèses, actes, Notez bien (Calibri)
NAVY = "143348"
TEAL = "008797"
TEXTE = "284D68"
GRIS = "657A88"
CLAIR = "F3F7FB"
ZEBRE = "E8F3F6"
BLANC = "FFFFFF"
ALERTE_T = "A96108"
ALERTE_F = "FFF2CC"

# Feuilles médecins / jours (Aptos, ancienne disposition)
NUIT = "123047"
TEAL2 = "087F8C"
TETE = "D4ECEE"
BLOC = "EAF4F5"
HORS = "EEF0F3"          # jours hors période
DIM = "EEE5F7"           # dimanche (cellules)
DIM_TETE = "6B538C"      # dimanche (en-tête)
ORANGE = "FCE4D6"        # au-dessus de la limite
FILET = "9ABBC2"

NB = "#,##0;[Red](#,##0);0"
PCT = "0.0%"
DEC = "0.00"
DATE = "dd/mm/yyyy"
NB_TIRET = "0;\\-0;\\–"


@lru_cache(maxsize=None)
def _font(nom, taille, gras, couleur, souligne=False):
    return Font(name=nom, size=taille, bold=gras, color=couleur, underline="single" if souligne else None)


@lru_cache(maxsize=None)
def _fill(couleur):
    return PatternFill("solid", fgColor=couleur) if couleur else PatternFill(fill_type=None)


@lru_cache(maxsize=None)
def _align(h, v, wrap, indent):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


@lru_cache(maxsize=None)
def _border_haut(couleur):
    return Border(top=Side(style="thin", color=couleur))


def style(cell, taille=10, gras=False, couleur=TEXTE, fond=None, fmt=None, h="left", v="center",
          wrap=True, indent=0, police="Calibri", filet=None, souligne=False):
    cell.font = _font(police, taille, gras, couleur, souligne)
    if fond is not None:
        cell.fill = _fill(fond)
    if fmt:
        cell.number_format = fmt
    cell.alignment = _align(h, v, wrap, indent)
    if filet:
        cell.border = _border_haut(filet)


def ecrire(ws, ref, valeur=None, **kw):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    if valeur is not None:
        c.value = valeur
    style(c, **kw)
    return c


def bandeau(ws, ligne, texte, derniere_col, taille=11, fond=NAVY, hauteur=26, couleur=BLANC, gras=True):
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=derniere_col)
    ecrire(ws, (ligne, 1), texte, taille=taille, gras=gras, couleur=couleur, fond=fond)
    ws.row_dimensions[ligne].height = hauteur


def note(ws, ligne, texte, derniere_col, alerte=False, taille=10, hauteur=34):
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=derniere_col)
    if alerte:
        ecrire(ws, (ligne, 1), texte, taille=taille, couleur=ALERTE_T, fond=ALERTE_F)
    else:
        ecrire(ws, (ligne, 1), texte, taille=taille, couleur=GRIS, fond=CLAIR)
    ws.row_dimensions[ligne].height = hauteur


def lien(ws, ref, texte, cible, fond=ZEBRE, couleur=TEAL, taille=10):
    c = ecrire(ws, ref, texte, taille=taille, gras=True, couleur=couleur, fond=fond)
    c.hyperlink = f"#{cible}"
    return c


def mise_en_page(ws, onglet, zoom=85, figer=None, resume_dessous=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.tabColor = onglet
    ws.sheet_properties.outlinePr.summaryBelow = resume_dessous
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if figer:
        ws.freeze_panes = figer
