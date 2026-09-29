"""Historique local des périodes traitées (aucune donnée patient : uniquement des totaux)."""
import datetime as dt
import os

import openpyxl

from . import visites as V
from .calculs import indicateurs


def _lire(chemin):
    if not os.path.exists(chemin):
        return []
    wb = openpyxl.load_workbook(chemin, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    entete = list(rows[0])
    return [dict(zip(entete, r)) for r in rows[1:] if r and r[0]]


def _ecrire(chemin, lignes):
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    cles = []
    for l in lignes:
        for k in l:
            if k not in cles:
                cles.append(k)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historique"
    ws.append(cles)
    for l in lignes:
        ws.append([l.get(k) for k in cles])
    wb.save(chemin)


def mettre_a_jour(chemin, cle_periode, ligne):
    """Remplace ou ajoute la ligne de la période ; renvoie toutes les lignes triées par début de période."""
    try:
        lignes = [l for l in _lire(chemin) if l.get("Clé") != cle_periode]
    except Exception:       # fichier illisible : on repart d'un historique vide
        lignes = []
    ligne = dict(ligne, **{"Clé": cle_periode})
    lignes.append(ligne)
    lignes.sort(key=lambda l: (str(l.get("Clé"))))
    try:
        _ecrire(chemin, lignes)
    except OSError:
        print("   NOTE : historique non enregistré (fichier ouvert dans Excel ?).")
    return lignes


def ligne_capacites(R):
    cs = indicateurs(R, "CSMKL2", V.PRINCIPALE)
    amb = indicateurs(R, "CHME", V.AMBULATOIRE)
    urg = indicateurs(R, "CHME", V.URGENCES)
    n = R.ndays
    return {
        "Période": f"{R.debut:%d/%m/%Y} - {R.fin:%d/%m/%Y}", "Jours": n,
        "CSMKL2\nvisites / jour": cs["total"] / n, "CSMKL2\nutilisation": cs["util"],
        "CHME amb.\nvisites / jour": amb["total"] / n, "CHME amb.\nutilisation": amb["util"],
        "Urgences\nvisites / jour": urg["total"] / n, "Lits occupés\nmoy. / jour": R.lits_occ["moy"],
        "Occupation\ndes lits": R.lits_occ["taux"],
        "Actes GPS\npar jour": sum(1 for a in R.actes if a["logiciel"] == "GPS") / n,
        "Actes Evo\npar jour": sum(1 for a in R.actes if a["logiciel"] != "GPS") / n,
        "Capacité\nCHME": "Horaires RDV" if R.horaire else f"Repère {R.cap_jour}",
        "Calculé le": dt.datetime.now().replace(microsecond=0),
    }
