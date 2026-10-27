"""Classeur « CSMKL2 et CHME » produit en même temps que « Activités et capacités », organisé par type de tableau pour que
les colonnes aient la même largeur d'un site à l'autre :
- onglet « Charge individuelle » : CSMKL2 (activité principale, puis urgences CHME dessous) à gauche, CHME ambulatoire à
  droite (spécialistes en cabinet, puis intervenants hors cabinet), mêmes colonnes ;
- onglet « Lecture quotidienne » : CSMKL2, CHME ambulatoire et urgences CHME l'un sous l'autre, mêmes colonnes.
Les infirmiers d'Evolucare (nom terminé par INF) ne figurent pas dans les listes : ils ne consultent pas."""
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

from . import visites as V
from .calculs import hors_cabinet, non_consultant
from .centrage import centrer
from .feuilles_synthese import cellule, entete_tableau, fusion, section, tableau_lecture, titre_periode
from .styles import CLAIR, TEXTE, bandeau, ecrire, mise_en_page

ONGLET_CHARGE = "167D8D"
ONGLET_LECTURE = "143348"
COLS_CHARGE = ["Intervenant", "Spécialité", "Consultation", "Avant 1 sem.", "Après 1 sem.", "Visites", "Jours actifs",
               "Moy / jour"]
LARGEURS_CHARGE = [30, 26, 12, 12, 12, 11, 11, 11]
GAUCHE, DROITE = 1, len(COLS_CHARGE) + 2      # colonnes de départ des deux blocs (une colonne vide entre les deux)
DER_CHARGE = DROITE + len(COLS_CHARGE) - 1
DER_LECTURE = 14


def specialites_medecins(R, site, act):
    """Spécialité de chaque médecin = celle de la majorité de ses visites sur la période."""
    c = defaultdict(Counter)
    for v in R.visites:
        if v["site"] == site and v["activite"] == act and v["medecin"]:
            c[v["medecin"]][v["specialite"]] += 1
    return {m: x.most_common(1)[0][0] for m, x in c.items()}


def _onglet(wb, nom, onglet, titre, sous_titre, der):
    ws = wb.create_sheet(nom)
    mise_en_page(ws, onglet, zoom=85)
    ws.page_setup.fitToHeight = 0
    bandeau(ws, 1, titre, der, taille=16, hauteur=30)
    fusion(ws, 2, 1, der)
    ecrire(ws, (2, 1), sous_titre, taille=10, couleur=TEXTE, fond=CLAIR)
    ws.row_dimensions[2].height = 34
    return ws


def tableau_charge_simple(ws, r, c0, R, site, act, titre, specialites, garder=None):
    """CHARGE INDIVIDUELLE (colonnes retenues par le directeur des opérations) à partir de la colonne c0 : un intervenant
    par ligne, classé sur la moyenne de visites par jour actif. garder(médecin) filtre les intervenants. Renvoie la ligne
    suivante."""
    c1 = c0 + len(COLS_CHARGE) - 1
    section(ws, r, titre, c0, c1)
    entete_tableau(ws, r + 1, [(c0 + i, c0 + i, t) for i, t in enumerate(COLS_CHARGE)])
    r += 2
    par_med = defaultdict(list)
    for (_, m), c in R.medjour[(site, act)].items():
        if not non_consultant(R, m) and (garder is None or garder(m)):
            par_med[m].append(c)
    total = {m: sum(c["total"] for c in cs) for m, cs in par_med.items()}
    moyenne = {m: total[m] / len(cs) for m, cs in par_med.items()}
    if not par_med:
        cellule(ws, r, c0, "Aucun intervenant sur la période", h="left", fmt=None, c2=c1)
        return r + 1
    for m in sorted(par_med, key=lambda m: (-moyenne[m], -total[m], m)):
        cs = par_med[m]
        cellule(ws, r, c0, m, h="left", fmt=None)
        cellule(ws, r, c0 + 1, specialites.get(m, "–"), h="left", fmt=None)
        for k, v in enumerate([sum(c["cons"] for c in cs), sum(c["avant"] for c in cs), sum(c["apres"] for c in cs),
                               total[m], len(cs)], start=2):
            cellule(ws, r, c0 + k, v)
        cellule(ws, r, c0 + 7, moyenne[m], fmt="0.0", gras=True)
        ws.row_dimensions[r].height = 22
        r += 1
    return r


def _onglet_charge(wb, R, per):
    ws = _onglet(wb, "Charge individuelle", ONGLET_CHARGE, f"CSMKL2 ET CHME / CHARGE INDIVIDUELLE — {per}",
                 "GPS + Evolucare • un intervenant par ligne, classé sur la moyenne de visites par jour actif (visites / jours "
                 "avec au moins une visite) • à gauche : CSMKL2 activité principale, puis urgences CHME • à droite : CHME "
                 "ambulatoire, spécialistes en cabinet puis intervenants hors cabinet (kinésithérapie, soins, imagerie, "
                 "pharmacie, médecine générale…) • infirmiers (nom terminé par INF) retirés : ils ne consultent pas",
                 DER_CHARGE)
    for c0 in (GAUCHE, DROITE):
        for i, l in enumerate(LARGEURS_CHARGE):
            ws.column_dimensions[get_column_letter(c0 + i)].width = l
    ws.column_dimensions[get_column_letter(DROITE - 1)].width = 3

    # Gauche : CSMKL2, puis urgences CHME (même repère de capacité)
    r = tableau_charge_simple(ws, 4, GAUCHE, R, "CSMKL2", V.PRINCIPALE, "CSMKL2 / ACTIVITÉ PRINCIPALE",
                              specialites_medecins(R, "CSMKL2", V.PRINCIPALE))
    tableau_charge_simple(ws, r + 2, GAUCHE, R, "CHME", V.URGENCES, "CHME / URGENCES",
                          specialites_medecins(R, "CHME", V.URGENCES))

    # Droite : CHME ambulatoire, spécialistes en cabinet puis hors cabinet
    spec = specialites_medecins(R, "CHME", V.AMBULATOIRE)
    hors = lambda m: hors_cabinet(R, m, spec.get(m))
    r = tableau_charge_simple(ws, 4, DROITE, R, "CHME", V.AMBULATOIRE, "CHME / AMBULATOIRE — SPÉCIALISTES EN CABINET", spec,
                              garder=lambda m: not hors(m))
    tableau_charge_simple(ws, r + 2, DROITE, R, "CHME", V.AMBULATOIRE, "CHME / AMBULATOIRE — HORS CABINET DE SPÉCIALISTE", spec,
                          garder=hors)
    return ws


def _onglet_lecture(wb, R, per):
    seuil = f"cabinet compté à partir de {R.seuil} visites dans la journée"
    capacite = (f"créneaux de rendez-vous de chaque médecin, limités à sa plage de présence (à défaut : médiane de la "
                f"spécialité, puis {R.cap_jour})" if R.horaire else f"{R.cap_jour} visites par médecin en cabinet et par jour")
    ws = _onglet(wb, "Lecture quotidienne", ONGLET_LECTURE, f"CSMKL2 ET CHME / LECTURE QUOTIDIENNE — {per}",
                 f"GPS + Evolucare • {seuil} • CSMKL2 : {R.cap_jour} visites par médecin et par jour, "
                 f"{R.cabinets_csmkl2} cabinets physiques • CHME ambulatoire : capacité = {capacite} ; utilisation sur les "
                 f"seules visites des cabinets comptés ; {R.cabinets_chme} cabinets physiques, "
                 "un cabinet peut accueillir deux spécialistes qui se succèdent • urgences : "
                 + (f"{R.garde_urgences} médecin(s) de garde × {R.cap_jour} visites par jour" if R.garde_urgences
                    else f"{R.cap_jour} visites par médecin et par jour") + " • infirmiers (nom terminé par INF) : visites "
                 "gardées, sans cabinet ni capacité, hors utilisation", DER_LECTURE)
    ws.column_dimensions["A"].width = 16
    for c in "BCDEFGHIJK":
        ws.column_dimensions[c].width = 12
    for c in "LMN":
        ws.column_dimensions[c].width = 10
    r = tableau_lecture(ws, 4, R, "CSMKL2", V.PRINCIPALE, DER_LECTURE, "CSMKL2 / ACTIVITÉ PRINCIPALE",
                        f"Intervenants >{R.cap_jour}", "Cabinets toutes activités", R.cab_site["CSMKL2"], None, R.cabinets_csmkl2)
    lib_depasse = "Intervenants >\ncapacité" if R.horaire else f"Intervenants >{R.cap_jour}"
    r = tableau_lecture(ws, r + 2, R, "CHME", V.AMBULATOIRE, DER_LECTURE, "CHME / AMBULATOIRE", lib_depasse,
                        f"Cabinets occupés en\nmême temps (sur {R.cabinets_chme})", R.cab_chme_spec, R.cabinets_chme, None)
    if R.garde_urgences:
        tableau_lecture(ws, r + 2, R, "CHME", V.URGENCES, DER_LECTURE, "CHME / URGENCES", f"Intervenants >{R.cap_jour}",
                        "Médecins nommés", R.medecins_nommes[("CHME", V.URGENCES)], None, None, lib_cabinets="Médecins\nde garde")
    else:
        tableau_lecture(ws, r + 2, R, "CHME", V.URGENCES, DER_LECTURE, "CHME / URGENCES", f"Intervenants >{R.cap_jour}",
                        "Cabinets toutes activités", R.cab_site["CHME"], None, None)
    return ws


def construire(R, chemin):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    per = titre_periode(R)
    _onglet_charge(wb, R, per)
    _onglet_lecture(wb, R, per)
    centrer(wb)
    wb.save(chemin)
    return wb
