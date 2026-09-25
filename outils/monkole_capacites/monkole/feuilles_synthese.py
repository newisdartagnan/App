"""Feuilles de synthèse : Dashboard, CSMKL2, CHME."""
from collections import defaultdict

from openpyxl.chart import LineChart, Reference, Series

from . import visites as V
from .calculs import indicateurs
from .feuilles_detail import libelle_periode
from .styles import (ALERTE_F, ALERTE_T, BLANC, CLAIR, DATE, DEC, GRIS, NAVY, NB, PCT, TEAL, TEXTE, ZEBRE, bandeau, ecrire,
                     lien, mise_en_page)
from .texte import MOIS, fr_nombre_espace as fr, fr_pct, nom_jour

LAB = "Laboratoire"
IMAGERIE = "Imagerie Médicale"
DASH_SPECS = ["Laboratoire", "Imagerie Médicale", "Nursing", "Hospitalisation", "Chirurgie", "Pédiatrie", "Ophtalmologie",
              "Dentisterie", "ORL"]


def zebre(r):
    return ZEBRE if r % 2 == 0 else BLANC


def fusion(ws, r, c1, c2, r2=None):
    if c2 > c1 or (r2 and r2 > r):
        ws.merge_cells(start_row=r, start_column=c1, end_row=r2 or r, end_column=c2)


def entete_tableau(ws, r, colonnes, hauteur=36, taille=10):
    """colonnes : liste de (col_debut, col_fin, texte)."""
    for c1, c2, t in colonnes:
        fusion(ws, r, c1, c2)
        for c in range(c1, c2 + 1):
            ecrire(ws, (r, c), t if c == c1 else None, taille=taille, gras=True, couleur=BLANC, fond=NAVY)
    ws.row_dimensions[r].height = hauteur


def cellule(ws, r, c1, valeur, fmt=NB, c2=None, h="right", fond=None, gras=False, couleur=TEXTE, taille=10):
    fond = fond or zebre(r)
    if c2:
        fusion(ws, r, c1, c2)
        for c in range(c1 + 1, c2 + 1):
            ecrire(ws, (r, c), taille=taille, gras=gras, couleur=couleur, fond=fond, fmt=fmt, h=h)
    if isinstance(valeur, str) and fmt in (NB, PCT, DEC):
        fmt = None
    return ecrire(ws, (r, c1), valeur, taille=taille, gras=gras, couleur=couleur, fond=fond, fmt=fmt, h=h)


def kpi(ws, c1, c2, titre, valeur, commentaire, fmt=NB, taille=27, lien_cible=None, lien_texte=None):
    if lien_texte:
        fusion(ws, 3, c1, c2)
        lien(ws, (3, c1), lien_texte, lien_cible)
        for c in range(c1 + 1, c2 + 1):
            ecrire(ws, (3, c), fond=ZEBRE)
    fusion(ws, 4, c1, c2)
    for c in range(c1, c2 + 1):
        ecrire(ws, (4, c), titre if c == c1 else None, gras=True, couleur=BLANC, fond=TEAL)
    fusion(ws, 5, c1, c2, r2=6)
    ecrire(ws, (5, c1), valeur, taille=taille, gras=True, couleur=TEAL, fond=BLANC, fmt=fmt if not isinstance(valeur, str) else None,
           h="center", wrap=None)
    fusion(ws, 7, c1, c2)
    for c in range(c1, c2 + 1):
        ecrire(ws, (7, c), commentaire if c == c1 else None, taille=9, couleur=GRIS, fond=CLAIR, h="center")


def haut_de_page(ws, titre, sous_titre, der):
    bandeau(ws, 1, titre, der, taille=19, hauteur=32)
    fusion(ws, 2, 1, der)
    ecrire(ws, "A2", sous_titre, fond=CLAIR)
    for r, h in ((2, 28), (3, 24), (4, 26), (5, 28), (6, 15), (7, 31)):
        ws.row_dimensions[r].height = h


def section(ws, r, texte, c1, c2, taille=11, hauteur=26):
    fusion(ws, r, c1, c2)
    for c in range(c1, c2 + 1):
        ecrire(ws, (r, c), texte if c == c1 else None, taille=taille, gras=True, couleur=BLANC, fond=NAVY)
    ws.row_dimensions[r].height = hauteur


def note_bloc(ws, r, texte, c1, c2, alerte=False, taille=10, hauteur=None):
    fusion(ws, r, c1, c2)
    for c in range(c1, c2 + 1):
        ecrire(ws, (r, c), texte if c == c1 else None, taille=taille, couleur=ALERTE_T if alerte else GRIS,
               fond=ALERTE_F if alerte else CLAIR)
    if hauteur:
        ws.row_dimensions[r].height = hauteur


def graphique(ws, ancre, source, premiere_ligne, n, series, col_etiquette, titre, axe="Visites"):
    ch = LineChart()
    ch.title = titre
    ch.y_axis.title = axe
    ch.height = 7.5
    ch.width = 15
    ch.legend.position = "b"
    for col, nom in series:
        data = Reference(source, min_col=col, min_row=premiere_ligne, max_row=premiere_ligne + n - 1)
        s = Series(data, title=nom)
        s.smooth = False
        ch.series.append(s)
    ch.set_categories(Reference(source, min_col=col_etiquette, min_row=premiere_ligne, max_row=premiere_ligne + n - 1))
    ws.add_chart(ch, ancre)


def titre_periode(R):
    return f"{libelle_periode(R, '-')} {R.fin.year}"


# ----------------------------------------------------------------------------------------------
# CSMKL2
# ----------------------------------------------------------------------------------------------
def feuille_csmkl2(wb, R, positions_actes, arbre_actes, onglet):
    ws = wb.create_sheet("CSMKL2")
    mise_en_page(ws, onglet, zoom=85, figer="A4")
    der = 12
    ws.column_dimensions["A"].width = 30
    for c in "BCDEFGHIJKL":
        ws.column_dimensions[c].width = 12
    ind = indicateurs(R, "CSMKL2", V.PRINCIPALE)
    lignes = R.jours_act[("CSMKL2", V.PRINCIPALE)]
    jours_plus = sum(1 for l in lignes if l["cab"] > R.cabinets_csmkl2)
    haut_de_page(ws, "CSMKL2 / CAPACITÉ DES ACTIVITÉS PRINCIPALES",
                 f"{titre_periode(R)} • GPS + Evolucare • MAISON ROSE et actes conservés en complément", der)
    kpi(ws, 1, 3, "CAPACITÉ UTILISÉE", ind["util"] if ind["util"] is not None else "N/D", f"sur {ind['cap']:,} visites théoriques",
        fmt=PCT, lien_cible="'Dashboard'!A1", lien_texte="Dashboard")
    kpi(ws, 4, 6, "VISITES PRINCIPALES", ind["total"], "Toutes les visites principales, isolées incluses",
        lien_cible="'CSMKL2 médecins'!A1", lien_texte="Médecins par jour")
    kpi(ws, 7, 9, "JOURS AU-DELÀ CAPACITÉ", ind["jour_audela"], f"{R.cap_jour} × cabinets principaux comptés / jour",
        lien_cible="'CSMKL2 jours'!A1", lien_texte="Jours par médecin")
    kpi(ws, 10, 12, f"JOURS AVEC PLUS DE {R.cabinets_csmkl2} CABINETS", jours_plus,
        f"Vérifier le partage et les rotations, pas {R.cabinets_csmkl2 + 1} salles supposées",
        lien_cible="'CSMKL2 actes'!A1", lien_texte="Actes détaillés")
    section(ws, 9, "VISITES PRINCIPALES ET CAPACITÉ DU MODÈLE", 1, der)
    for r in range(10, 27):
        ws.row_dimensions[r].height = 20
    graphique(ws, "A10", wb["_Jours"], R.lignes_jours[("CSMKL2", V.PRINCIPALE)], R.ndays, [(5, "Visites"), (12, "Capacité")], 20,
              titre_periode(R).replace(f" {R.fin.year}", ""))
    dp = ind["date_pic"]
    note_bloc(ws, 27, f"Pic : {ind['pic_jour']} visites le {dp.strftime('%d/%m') if dp else '–'} = {ind['gps_pic']} GPS + "
                      f"{ind['evo_pic']} Evolucare. {ind['isolees']} visites principales isolées et {ind['sans']} sans médecin restent comptées.",
              1, der, hauteur=34)
    # Lecture quotidienne
    section(ws, 29, "LECTURE QUOTIDIENNE / ACTIVITÉ PRINCIPALE", 1, der)
    cols = ["Jour", "GPS", "Evolucare", "Visites réelles", "Cabinets", "Capacité", "Utilisation", "Marge cumulée",
            f"Intervenants >{R.cap_jour}", "Sans médecin", "Visites isolées", "Cabinets toutes activités"]
    entete_tableau(ws, 30, [(i, i, t) for i, t in enumerate(cols, start=1)])
    r = 31
    for l in lignes:
        j = l["jour"]
        cellule(ws, r, 1, f"{nom_jour(j)} {j.strftime('%d/%m')}", h="left", fmt=None)
        for c, v in enumerate([l["gps"], l["evo"], l["total"], l["cab"], l["cap"]], start=2):
            cellule(ws, r, c, v)
        cellule(ws, r, 7, l["util"] if l["util"] is not None else "N/D", fmt=PCT)
        for c, v in enumerate([l["marge"], l["depass"], l["sans"], l["isolees"], R.cab_site["CSMKL2"][j]], start=8):
            cellule(ws, r, c, v)
        ws.row_dimensions[r].height = 26
        r += 1
    # Charge individuelle
    r += 2
    section(ws, r, "CHARGE INDIVIDUELLE / ACTIVITÉ PRINCIPALE", 1, der)
    r += 1
    cols = ["Intervenant", "Consultation", "Avant 1 sem.", "Après 1 sem.", "Visites", "Jours actifs", "Min / jour actif",
            "Max / jour", "GPS au pic", "Evo au pic", "Date du pic", f"Jours >{R.cap_jour}"]
    entete_tableau(ws, r, [(i, i, t) for i, t in enumerate(cols, start=1)])
    r += 1
    md = R.medjour[("CSMKL2", V.PRINCIPALE)]
    par_med = defaultdict(dict)
    for (j, m), c in md.items():
        par_med[m][j] = c
    ordre = sorted(par_med, key=lambda m: (min(par_med[m]), m))
    for m in ordre:
        jours_m = par_med[m]
        pic_j = min(jours_m, key=lambda j: (-jours_m[j]["total"], j))
        vals = [sum(c["cons"] for c in jours_m.values()), sum(c["avant"] for c in jours_m.values()),
                sum(c["apres"] for c in jours_m.values()), sum(c["total"] for c in jours_m.values()), len(jours_m),
                min(c["total"] for c in jours_m.values()), jours_m[pic_j]["total"], jours_m[pic_j]["gps"], jours_m[pic_j]["evo"]]
        cellule(ws, r, 1, m, h="left", fmt=None)
        for c, v in enumerate(vals, start=2):
            cellule(ws, r, c, v)
        cellule(ws, r, 11, pic_j, fmt=DATE)
        cellule(ws, r, 12, sum(c["audela"] for c in jours_m.values()))
        ws.row_dimensions[r].height = 30
        r += 1
    # MAISON ROSE
    r += 2
    section(ws, r, "MAISON ROSE / HORS CAPACITÉ PRINCIPALE", 1, der)
    r += 1
    entete_tableau(ws, r, [(1, 1, "Activité / UF"), (2, 2, "GPS"), (3, 3, "Evolucare"), (4, 4, "Total visites")])
    for c in range(5, der + 1):
        ecrire(ws, (r, c), gras=True, couleur=BLANC, fond=NAVY)
    r += 1
    mr = defaultdict(lambda: [0, 0])
    for v in R.visites:
        if v["site"] == "CSMKL2" and v["activite"] == V.SECONDAIRE:
            mr[v["specialite"]][0 if v["logiciel"] == "GPS" else 1] += 1
    for sp in sorted(mr, key=lambda s: s.lower()):
        g, e = mr[sp]
        cellule(ws, r, 1, sp, h="left", fmt=None)
        cellule(ws, r, 2, g)
        cellule(ws, r, 3, e)
        cellule(ws, r, 4, g + e)
        for c in range(5, der + 1):
            ecrire(ws, (r, c), fond=zebre(r), h="right")
        ws.row_dimensions[r].height = 24
        r += 1
    # Actes
    r += 2
    section(ws, r, "ACTES / PRESTATIONS DU SITE — SOURCES ET DATES DISTINGUÉES", 1, der)
    r += 1
    entete_tableau(ws, r, [(1, 1, "Spécialité"), (2, 2, "GPS / Date_V"), (3, 3, "Evo / DATEHEURE"), (4, 4, "Dossiers GPS"),
                           (5, 5, "Dossiers Evo")])
    for c in range(6, der + 1):
        ecrire(ws, (r, c), gras=True, couleur=BLANC, fond=NAVY)
    r += 1
    for sp in arbre_actes:
        a = sp["agg"]
        c = cellule(ws, r, 1, sp["nom"], h="left", fmt=None)
        c.hyperlink = f"#'CSMKL2 actes'!A{positions_actes[sp['nom']]}"
        for col, v in enumerate([a["gps"], a["evo"], a["dos_gps"], a["dos_evo"]], start=2):
            cellule(ws, r, col, v)
        for col in range(6, der + 1):
            ecrire(ws, (r, col), fond=zebre(r), h="right")
        ws.row_dimensions[r].height = 24
        r += 1
    r += 2
    note_bloc(ws, r, f"{ind['depass']} journées-intervenants dépassent {R.cap_jour} visites; pic individuel {ind['pic_medecin']}. "
                     f"Examiner les répartitions et horaires avant toute décision de renfort. Les {R.cabinets_csmkl2} cabinets restent "
                     f"l’hypothèse physique du modèle.", 1, der, alerte=True, hauteur=44)
    return ws


# ----------------------------------------------------------------------------------------------
# CHME
# ----------------------------------------------------------------------------------------------
def actes_lies(R, specialite, activite=None):
    cles = {v["cle"] for v in R.visites if v["site"] == "CHME" and v["specialite"] == specialite
            and (activite is None or v["activite"] == activite)}
    g = sum(1 for a in R.actes if a["cle"] in cles and a["logiciel"] == "GPS")
    e = sum(1 for a in R.actes if a["cle"] in cles and a["logiciel"] != "GPS")
    return g, e


def feuille_chme(wb, R, onglet):
    ws = wb.create_sheet("CHME")
    mise_en_page(ws, onglet, zoom=85, figer="A4")
    der = 12
    ws.column_dimensions["A"].width = 32
    for c in "BCDEFGHIJKL":
        ws.column_dimensions[c].width = 12
    t = R.unites_total
    sans = t["sans_gps"] + t["sans_evo"]
    haut_de_page(ws, "CHME / CAPACITÉ ET ACTIVITÉS",
                 f"{titre_periode(R)} • GPS + Evolucare • Visites, actes et hospitalisation; sources distinguées", der)
    kpi(ws, 1, 3, "OCCUPATION DOCUMENTÉE / GPS", t["occ_gps"], f"Huit unités / {R.total_lits} lits / durées GPS connues", fmt=PCT,
        lien_cible="'Dashboard'!A1", lien_texte="Dashboard")
    kpi(ws, 4, 6, "OCCUPATION RECONSTITUÉE / EVO", t["occ_evo"], "Dates sans heures et sorties futures : à confirmer", fmt=PCT,
        lien_cible="'CHME médecins'!A1", lien_texte="Médecins par jour")
    kpi(ws, 7, 9, "DOSSIERS HOSPI / GPS - EVO", f"{t['dos_gps']} / {t['dos_evo']}", "Pas de fusion de patients entre les logiciels",
        lien_cible="'CHME jours'!A1", lien_texte="Jours et séjours")
    kpi(ws, 10, 12, "DOSSIERS SANS DATE D’ENTRÉE", sans, "Comptés; aucune durée inventée",
        lien_cible="'CHME actes'!A1", lien_texte="Actes détaillés")
    section(ws, 9, "HOSPITALISATION / OCCUPATION PAR SOURCE — NE PAS ADDITIONNER LES COURBES", 1, der)
    for r in range(10, 27):
        ws.row_dimensions[r].height = 20
    graphique(ws, "A10", wb["_Hospi jour"], 6, R.ndays, [(2, "GPS journées / 8 unités"), (3, "Evo journées / 8 unités"),
                                                          (4, "Lits CHME")], 10, "Moyenne sur 24 heures, par source", axe="Lits occupés")
    note_bloc(ws, 27, "Les taux sont partiels : séjours sans entrée exclus des durées et éventuels séjours communs GPS/Evo non "
                      "rapprochables. Le complément à 100 % ne prouve pas des lits libres.", 1, der, alerte=True, hauteur=34)
    section(ws, 29, "HOSPITALISATION / RÉPARTITION PAR UNITÉ", 1, der)
    cols = ["Unité", "Lits", "Dossiers\nGPS", "Dossiers\nEvolucare", "Sans entrée\nGPS", "Sans entrée\nEvolucare", "Jours\nGPS",
            "Jours\nEvolucare", "Occupation\nGPS", "Occupation\nEvolucare", "Sorties\nGPS", "Sorties\nEvolucare"]
    entete_tableau(ws, 30, [(i, i, x) for i, x in enumerate(cols, start=1)])
    r = 31
    for u in R.unites:
        cellule(ws, r, 1, u["unite"], h="left", fmt=None)
        cellule(ws, r, 2, u["lits"] if u["lits"] else "N/D")
        for c, v in enumerate([u["dos_gps"], u["dos_evo"], u["sans_gps"], u["sans_evo"]], start=3):
            cellule(ws, r, c, v)
        cellule(ws, r, 7, u["jours_gps"], fmt=DEC)
        cellule(ws, r, 8, u["jours_evo"], fmt=DEC)
        cellule(ws, r, 9, u["occ_gps"] if u["lits"] else "Hors capacité", fmt=PCT)
        cellule(ws, r, 10, u["occ_evo"] if u["lits"] else "Hors capacité", fmt=PCT)
        cellule(ws, r, 11, u["sorties_gps"])
        cellule(ws, r, 12, u["sorties_evo"])
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1
    tot = [t["lits"], t["dos_gps"], t["dos_evo"], t["sans_gps"], t["sans_evo"]]
    ecrire(ws, (r, 1), "TOTAL CHME", gras=True, couleur=BLANC, fond=TEAL)
    for c, v in enumerate(tot, start=2):
        ecrire(ws, (r, c), v, gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    ecrire(ws, (r, 7), t["jours_gps"], gras=True, couleur=BLANC, fond=TEAL, fmt=DEC, h="right")
    ecrire(ws, (r, 8), t["jours_evo"], gras=True, couleur=BLANC, fond=TEAL, fmt=DEC, h="right")
    ecrire(ws, (r, 9), t["occ_gps"], gras=True, couleur=BLANC, fond=TEAL, fmt=PCT, h="right")
    ecrire(ws, (r, 10), t["occ_evo"], gras=True, couleur=BLANC, fond=TEAL, fmt=PCT, h="right")
    ecrire(ws, (r, 11), t["sorties_gps"], gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    ecrire(ws, (r, 12), t["sorties_evo"], gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    ws.row_dimensions[r].height = 28
    r += 1
    fusion(ws, r, 1, der)
    lien(ws, (r, 1), "Ouvrir la liste des dossiers hospitaliers, dates et durées", "'CHME jours'!A1")
    ws.row_dimensions[r].height = 24
    # Ambulatoire et urgences
    r += 2
    section(ws, r, "AMBULATOIRE ET URGENCES / VISITES ET CAPACITÉ", 1, der)
    r += 1
    cols = ["Activité", "GPS", "Evolucare", "Total visites", "Consultation", "Avant 1 sem.", "Après 1 sem.", "Cabinets-jours",
            "Capacité", "Utilisation", "Sans médecin", "Pic / jour"]
    entete_tableau(ws, r, [(i, i, x) for i, x in enumerate(cols, start=1)])
    r += 1
    for act in (V.AMBULATOIRE, V.URGENCES):
        ind = indicateurs(R, "CHME", act)
        cellule(ws, r, 1, act, h="left", fmt=None)
        for c, v in enumerate([ind["gps"], ind["evo"], ind["total"], ind["cons"], ind["avant"], ind["apres"], ind["cab"], ind["cap"]],
                              start=2):
            cellule(ws, r, c, v)
        cellule(ws, r, 10, ind["util"] if ind["util"] is not None else "N/D", fmt=PCT)
        cellule(ws, r, 11, ind["sans"])
        cellule(ws, r, 12, ind["pic_jour"])
        ws.row_dimensions[r].height = 30
        r += 1
    urg = indicateurs(R, "CHME", V.URGENCES)
    note_bloc(ws, r, f"Urgences : {urg['sans']} visites sans médecin renseigné restent dans le volume, sans ajouter de cabinet. "
                     "Le quotient affiché n’est pas une mesure fiable de la capacité réelle des urgences.", 1, der, alerte=True, hauteur=36)
    r += 2
    section(ws, r, "AMBULATOIRE / ÉVOLUTION DE LA CHARGE", 1, der)
    graphique(ws, f"A{r + 1}", wb["_Jours"], R.lignes_jours[("CHME", V.AMBULATOIRE)], R.ndays, [(5, "Visites"), (12, "Capacité")], 20,
              titre_periode(R).replace(f" {R.fin.year}", ""))
    for x in range(r + 1, r + 18):
        ws.row_dimensions[x].height = 20
    r += 19
    section(ws, r, "CONSULTATIONS PAR SPÉCIALITÉ / ACTIVITÉ ET ACTES DES DOSSIERS VUS", 1, der)
    r += 1
    cols = ["Spécialité / activité", "Consultation", "Avant 1 sem.", "Après 1 sem.", "Visites", "GPS", "Evolucare",
            f"Journées intervenants ≥{R.seuil}", "Repère visites", "Utilisation repère", "Actes liés GPS", "Actes liés Evo"]
    entete_tableau(ws, r, [(i, i, x) for i, x in enumerate(cols, start=1)])
    r += 1
    specs = defaultdict(list)
    for v in R.visites:
        if v["site"] == "CHME" and v["activite"] == V.AMBULATOIRE:
            specs[v["specialite"]].append(v)
    ordre = sorted(specs, key=lambda s: (-len(specs[s]), s))
    urg_vs = [v for v in R.visites if v["site"] == "CHME" and v["activite"] == V.URGENCES]
    for sp, vs, act in [(s, specs[s], V.AMBULATOIRE) for s in ordre] + [("Urgences générales", urg_vs, V.URGENCES)]:
        if not vs:
            continue
        cnt = defaultdict(int)
        for v in vs:
            if v["medecin"]:
                cnt[(v["jour"], v["medecin"])] += 1
        jint = sum(1 for n in cnt.values() if n >= R.seuil)
        rep = jint * R.cap_jour
        g, e = actes_lies(R, sp, act)
        vals = [sum(1 for v in vs if v["classe"] == V.CONSULTATION), sum(1 for v in vs if v["classe"] == V.AVANT),
                sum(1 for v in vs if v["classe"] == V.APRES), len(vs), sum(1 for v in vs if v["logiciel"] == "GPS"),
                sum(1 for v in vs if v["logiciel"] != "GPS"), jint, rep]
        cellule(ws, r, 1, sp, h="left", fmt=None)
        for c, v in enumerate(vals, start=2):
            cellule(ws, r, c, v)
        cellule(ws, r, 10, (len(vs) / rep) if rep else "N/D", fmt=PCT)
        cellule(ws, r, 11, g)
        cellule(ws, r, 12, e)
        ws.row_dimensions[r].height = 32
        r += 1
    r += 1
    note_bloc(ws, r, "Actes liés = tous les actes des dossiers vus dans cette spécialité, pas uniquement les actes produits par cette "
                     "spécialité. Un dossier multispecialités peut apparaître sur plusieurs lignes : ne pas les additionner.", 1, der,
              hauteur=40)
    r += 2
    note_bloc(ws, r, "Les repères par spécialité ne s’additionnent pas en capacité physique : un intervenant peut couvrir plusieurs "
                     "spécialités le même jour. Les soins, examens et pharmacie restent inclus selon le périmètre du modèle.", 1, der,
              hauteur=38)
    return ws


# ----------------------------------------------------------------------------------------------
# DASHBOARD
# ----------------------------------------------------------------------------------------------
def feuille_dashboard(wb, R, pos_chme, arbre_chme, arbre_cs, onglet):
    ws = wb.create_sheet("Dashboard", 0)
    mise_en_page(ws, onglet, zoom=80, figer="A4")
    der = 15
    for i in range(1, der + 1):
        ws.column_dimensions[chr(64 + i)].width = 12
    cs = indicateurs(R, "CSMKL2", V.PRINCIPALE)
    mr = indicateurs(R, "CSMKL2", V.SECONDAIRE)
    amb = indicateurs(R, "CHME", V.AMBULATOIRE)
    urg = indicateurs(R, "CHME", V.URGENCES)
    t = R.unites_total
    tot_cs = cs["total"] + mr["total"]
    tot_ch = amb["total"] + urg["total"]
    g_tot = sum(1 for a in R.actes if a["logiciel"] == "GPS")
    e_tot = sum(1 for a in R.actes if a["logiciel"] != "GPS")
    haut_de_page(ws, "MONKOLE / ACTIVITÉS ET CAPACITÉS PAR SITE",
                 f"Direction des opérations • {titre_periode(R)} • CSMKL2 + CHME • GPS + Evolucare", der)
    kpi(ws, 1, 4, "VISITES HORS HOSPITALISATION", tot_cs + tot_ch, f"CSMKL2 : {fr(tot_cs)} | CHME : {fr(tot_ch)}",
        lien_cible="'CSMKL2'!A1", lien_texte="CSMKL2 / modèle actualisé")
    kpi(ws, 5, 8, "ACTES / GPS — EVOLUCARE", f"{g_tot:,} / {e_tot:,}", "Date_V GPS / DATEHEURE Evo; séries séparées", taille=23,
        lien_cible="'CHME'!A1", lien_texte="CHME / modèle actualisé")
    kpi(ws, 9, 12, "DOSSIERS HOSPI CHME / GPS — EVO", f"{t['dos_gps']} / {t['dos_evo']}", "Dossiers par logiciel, pas patients uniques",
        lien_cible="'CHME actes'!A1", lien_texte="Actes par spécialité")
    kpi(ws, 13, 15, f"LITS CHME / {len(R.lits)} UNITÉS", R.total_lits, "Capacité reprise du modèle fourni",
        lien_cible="'Notez bien'!A1", lien_texte="Règles / contrôles")
    note_bloc(ws, 9, "Visites, actes et séjours sont trois mesures différentes. Ne pas les additionner. Les deux logiciels restent "
                     "identifiables; aucun dédoublonnage patient interlogiciels n’est présumé.", 1, der, hauteur=30)
    section(ws, 11, "01 / ACTIVITÉ ET CAPACITÉ DES CONSULTATIONS / LECTURE PAR SITE", 1, der)
    entete_tableau(ws, 12, [(1, 3, "Site / activité"), (4, 4, "GPS"), (5, 5, "Evolucare"), (6, 6, "Visites"), (7, 8, "Cabinets-jours"),
                            (9, 10, "Capacité du modèle"), (11, 12, "Utilisation"), (13, 15, "Lecture opérationnelle")], taille=9, hauteur=35)
    lignes = [
        ("CSMKL2 / principale", cs, f"{cs['depass']} journées-intervenants >{R.cap_jour}; pic à {cs['pic_medecin']}."),
        ("CSMKL2 / MAISON ROSE", mr, "Information; hors capacité principale."),
        ("CHME / ambulatoire", amb, f"{amb['depass']} journées-intervenants >{R.cap_jour}; pic à {amb['pic_medecin']}."),
        ("CHME / urgences", urg, f"{urg['sans']} visites sans médecin; repère incomplet."),
    ]
    for i, (nom, d, lecture) in enumerate(lignes):
        r = 13 + i
        cellule(ws, r, 1, nom, c2=3, h="left", fmt=None)
        cellule(ws, r, 4, d["gps"])
        cellule(ws, r, 5, d["evo"])
        cellule(ws, r, 6, d["total"])
        if nom.endswith("MAISON ROSE"):
            cellule(ws, r, 7, "Hors repère", c2=8, fmt=None)
            cellule(ws, r, 9, "Hors repère", c2=10, fmt=None)
            cellule(ws, r, 11, "N/D", c2=12, fmt=None)
        else:
            cellule(ws, r, 7, d["cab"], c2=8)
            cellule(ws, r, 9, d["cap"], c2=10)
            cellule(ws, r, 11, d["util"] if d["util"] is not None else "N/D", c2=12, fmt=PCT)
        cellule(ws, r, 13, lecture, c2=15, fmt=None)
        ws.row_dimensions[r].height = 35
    note_bloc(ws, 18, f"Repère conservé : {R.cap_jour} × intervenants comptés à partir de {R.seuil} visites/jour. Les visites isolées "
                      "restent au numérateur. Ce n’est pas un inventaire des cabinets physiques du CHME.", 1, der, hauteur=30)
    section(ws, 20, "CSMKL2 / VISITES PRINCIPALES ET CAPACITÉ", 1, 7, taille=10, hauteur=25)
    section(ws, 20, "CHME / VISITES AMBULATOIRES ET CAPACITÉ", 9, 15, taille=10, hauteur=25)
    for r in range(21, 37):
        ws.row_dimensions[r].height = 18
    per = titre_periode(R).replace(f" {R.fin.year}", "")
    graphique(ws, "A21", wb["_Jours"], R.lignes_jours[("CSMKL2", V.PRINCIPALE)], R.ndays, [(5, "Visites"), (12, "Capacité")], 20, per)
    graphique(ws, "I21", wb["_Jours"], R.lignes_jours[("CHME", V.AMBULATOIRE)], R.ndays, [(5, "Visites"), (12, "Capacité")], 20, per)
    jours_plus = sum(1 for l in R.jours_act[("CSMKL2", V.PRINCIPALE)] if l["cab"] > R.cabinets_csmkl2)
    note_bloc(ws, 37, f"CSMKL2 : {jours_plus} journée(s) avec plus de {R.cabinets_csmkl2} cabinets comptés. Vérifier les rotations.",
              1, 7, alerte=True, taille=9)
    note_bloc(ws, 37, "CHME : les capacités des appareils et les horaires des spécialités ne sont pas fournis.", 9, 15, taille=9)
    ws.row_dimensions[37].height = 32
    # 02 actes CHME / 03 hospitalisation
    section(ws, 39, "02 / CHME / ACTES ET PRESTATIONS", 1, 7, taille=10, hauteur=25)
    section(ws, 39, "03 / CHME / HOSPITALISATION", 9, 15, taille=10, hauteur=25)
    entete_tableau(ws, 40, [(1, 4, "Spécialité"), (5, 6, "GPS / Date_V"), (7, 7, "Evo / DATEHEURE"), (9, 11, "Unité"), (12, 12, "Lits"),
                            (13, 13, "Taux GPS"), (14, 14, "Taux Evo"), (15, 15, "Sans entrée")], taille=9, hauteur=35)
    par_nom = {sp["nom"]: sp for sp in arbre_chme}
    for i, nom in enumerate(DASH_SPECS):
        r = 41 + i
        sp = par_nom.get(nom)
        c = cellule(ws, r, 1, nom, c2=4, h="left", fmt=None, taille=11)
        if sp:
            c.hyperlink = f"#'CHME actes'!A{pos_chme[nom]}"
        cellule(ws, r, 5, sp["agg"]["gps"] if sp else 0, c2=6, taille=11, fond=BLANC)
        cellule(ws, r, 7, sp["agg"]["evo"] if sp else 0, taille=11, fond=BLANC)
    autres = [sp for sp in arbre_chme if sp["nom"] not in DASH_SPECS]
    r = 41 + len(DASH_SPECS)
    c = cellule(ws, r, 1, "Autres spécialités — voir détail +", c2=4, h="left", fmt=None, taille=11)
    c.hyperlink = "#'CHME actes'!A1"
    cellule(ws, r, 5, sum(sp["agg"]["gps"] for sp in autres), c2=6, taille=11, fond=BLANC)
    cellule(ws, r, 7, sum(sp["agg"]["evo"] for sp in autres), taille=11, fond=BLANC)
    for i, u in enumerate(R.unites):
        r = 41 + i
        cellule(ws, r, 9, u["unite"], c2=11, h="left", fmt=None)
        cellule(ws, r, 12, u["lits"] if u["lits"] else "N/D", taille=11, fond=BLANC)
        cellule(ws, r, 13, u["occ_gps"] if u["lits"] else "Hors capacité", fmt=PCT, taille=11, fond=BLANC)
        cellule(ws, r, 14, u["occ_evo"] if u["lits"] else "Hors capacité", fmt=PCT, taille=11, fond=BLANC)
        cellule(ws, r, 15, u["sans_gps"] + u["sans_evo"], taille=11, fond=BLANC)
        ws.row_dimensions[r].height = 29
    ws.row_dimensions[50].height = 28
    r = 51
    fusion(ws, r, 1, 4)
    for c in range(1, 5):
        ecrire(ws, (r, c), "TOTAL CHME / SOURCES SÉPARÉES" if c == 1 else None, gras=True, couleur=BLANC, fond=TEAL)
    fusion(ws, r, 5, 6)
    for c, v in ((5, sum(1 for a in R.actes if a["logiciel"] == "GPS" and a["site"] == "CHME")), (6, None),
                 (7, sum(1 for a in R.actes if a["logiciel"] != "GPS" and a["site"] == "CHME"))):
        ecrire(ws, (r, c), v, gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    fusion(ws, r, 9, 11)
    for c in range(9, 12):
        nb_u = "HUIT" if len(R.lits) == 8 else str(len(R.lits))
        ecrire(ws, (r, c), f"TOTAL / {nb_u} UNITÉS POUR LES TAUX" if c == 9 else None,
               taille=9, gras=True, couleur=BLANC, fond=TEAL)
    ecrire(ws, (r, 12), t["lits"], gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    ecrire(ws, (r, 13), t["occ_gps"], gras=True, couleur=BLANC, fond=TEAL, fmt=PCT, h="right")
    ecrire(ws, (r, 14), t["occ_evo"], gras=True, couleur=BLANC, fond=TEAL, fmt=PCT, h="right")
    ecrire(ws, (r, 15), t["sans_gps"] + t["sans_evo"], gras=True, couleur=BLANC, fond=TEAL, fmt=NB, h="right")
    ws.row_dimensions[r].height = 30
    note_bloc(ws, 53, "Evo : Date_V absente; DATEHEURE affichée à part. Quantités et produits dans le détail, sans montant.", 1, 7,
              alerte=True, taille=9)
    note_bloc(ws, 53, "Taux partiels, non additionnables entre logiciels. Dates de sortie Evo futures à confirmer; pas des lits libres "
                      "déduits.", 9, 15, alerte=True, taille=9)
    ws.row_dimensions[53].height = 35
    # 04 laboratoire / 05 imagerie
    section(ws, 55, "04 / LABORATOIRE / PRINCIPALES SOUS-SPÉCIALITÉS", 1, 7, taille=10, hauteur=26)
    section(ws, 55, "05 / IMAGERIE CHME / SOUS-SPÉCIALITÉS", 9, 15, taille=10, hauteur=26)
    entete_tableau(ws, 56, [(1, 4, "Laboratoire"), (5, 6, "GPS / Date_V"), (7, 7, "Evo / DATEHEURE"), (9, 12, "Imagerie"),
                            (13, 14, "GPS / Date_V"), (15, 15, "Evo / DATEHEURE")], taille=9, hauteur=35)
    lab = par_nom.get(LAB, {"subs": []})["subs"]
    img = par_nom.get(IMAGERIE, {"subs": []})["subs"]
    for i, sb in enumerate(lab):
        r = 57 + i
        cellule(ws, r, 1, sb["nom"], c2=4, h="left", fmt=None, taille=9)
        cellule(ws, r, 5, sb["agg"]["gps"], c2=6, taille=11, fond=BLANC)
        cellule(ws, r, 7, sb["agg"]["evo"], taille=11, fond=BLANC)
        ws.row_dimensions[r].height = 26
    for i, sb in enumerate(img):
        r = 57 + i
        cellule(ws, r, 9, sb["nom"], c2=12, h="left", fmt=None, taille=9)
        cellule(ws, r, 13, sb["agg"]["gps"], c2=14, taille=11, fond=BLANC)
        cellule(ws, r, 15, sb["agg"]["evo"], taille=11, fond=BLANC)
        ws.row_dimensions[r].height = 26
    rb = max(57 + len(img) + 2, 65)
    section(ws, rb, "CSMKL2 / ACTES PAR SOURCE", 9, 15, taille=10, hauteur=26)
    g_cs = sum(1 for a in R.actes if a["logiciel"] == "GPS" and a["site"] == "CSMKL2")
    e_cs = sum(1 for a in R.actes if a["logiciel"] != "GPS" and a["site"] == "CSMKL2")
    for k, (lib, v) in enumerate((("GPS / Date_V", g_cs), ("Evolucare / DATEHEURE", e_cs))):
        r = rb + 1 + k
        cellule(ws, r, 9, lib, c2=12, h="left", fmt=None, taille=9)
        cellule(ws, r, 13, v, taille=11, fond=BLANC)
        ws.row_dimensions[r].height = 26
    r = rb + 3
    fusion(ws, r, 9, 15)
    lien(ws, (r, 9), "Ouvrir les actes CSMKL2", "'CSMKL2 actes'!A1")
    ws.row_dimensions[r].height = 24
    # 06 points d'attention
    r0 = max(57 + len(lab), rb + 4) + 2
    section(ws, r0, "06 / POINTS D’ATTENTION POUR LE DIRECTEUR DES OPÉRATIONS", 1, der)
    hors = [b for b in R.base_hospi if b["logiciel"] == "Evolucare" and not b["dans_periode"]]
    mois_hors = sorted({MOIS[b["entree"].month - 1] for b in hors if b["entree"]})
    if len(hors) == 1 and mois_hors:
        txt_hors = f"1 entrée de {mois_hors[0]} est hors période."
    else:
        txt_hors = f"{len(hors)} entrée(s) hors période" + (f" ({', '.join(mois_hors)})." if mois_hors else ".")
    nb_sans_site = sum(1 for b in R.base_hospi if b["logiciel"] == "Evolucare" and b["dans_periode"] and b["site"] not in ("CHME", "CSMKL2"))
    points = [
        ("CSMKL2 / RÉPARTIR LA CHARGE", f"{cs['depass']} journées-intervenants dépassent {R.cap_jour} visites; {jours_plus} jour(s) avec plus "
                                        f"de {R.cabinets_csmkl2} cabinets comptés. Examiner les rotations avant de conclure à un besoin de renfort."),
        ("CHME / CHARGE PAR SPÉCIALITÉ", f"{amb['depass']} journées-intervenants ambulatoires dépassent {R.cap_jour}, malgré un taux global de "
                                         f"{fr_pct(amb['util'] or 0)}. Lire le détail par médecin : la moyenne du site masque les pointes."),
        ("URGENCES / IDENTITÉS", f"{urg['sans']} visites sans intervenant renseigné. Compléter ce champ avant de tirer une conclusion de "
                                 "capacité aux urgences."),
        ("HOSPITALISATION / FIABILITÉ", f"{t['sans_gps'] + t['sans_evo']} dossiers CHME sans date d’entrée. Les taux GPS et Evo restent "
                                        "séparés; aucun patient unique commun n’est identifié entre logiciels."),
        ("ACTES / DATES ET PÉRIMÈTRE", f"Date_V est disponible uniquement pour GPS. Les {fr(e_tot)} prestations Evo sont visibles séparément "
                                       f"sur DATEHEURE; {fr(len(R.produits))} lignes de produits sont hors actes."),
        ("SOURCE EVOLUCARE / SÉJOURS", f"{txt_hors} {nb_sans_site} séjours couvrant {MOIS[R.debut.month - 1]} restent sans site; ils ne "
                                       "sont pas attribués au CHME par supposition."),
    ]
    for i, (lib, txt) in enumerate(points):
        r = r0 + 2 + i
        fond = zebre(r)
        fusion(ws, r, 1, 4)
        for c in range(1, 5):
            ecrire(ws, (r, c), lib if c == 1 else None, taille=9, gras=True, couleur=TEAL, fond=fond)
        fusion(ws, r, 5, der)
        for c in range(5, der + 1):
            ecrire(ws, (r, c), txt if c == 5 else None, couleur=TEXTE, fond=fond)
        ws.row_dimensions[r].height = 42
    r = r0 + 9
    note_bloc(ws, r, "Périmètre : CSMKL2 et CHME documentés dans ces exports. Les visites et actes sont des enregistrements; les séjours "
                     "peuvent être incomplets. Aucun montant, facture ou diagnostic analysé dans ce fichier.", 1, der, hauteur=34)
    return ws
