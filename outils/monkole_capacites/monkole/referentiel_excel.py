"""Référentiel modifiable dans Excel (Referentiel_Monkole.xlsx), créé au premier lancement."""
import copy
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

from . import referentiel as D

ENTETE = PatternFill("solid", fgColor="143348")


def defauts():
    return {k: copy.deepcopy(getattr(D, k)) for k in dir(D) if k.isupper()}


def _feuille(wb, nom, entetes, lignes, largeurs, aide):
    ws = wb.create_sheet(nom)
    ws.append([aide])
    ws["A1"].font = Font(italic=True, color="657A88")
    ws.append(entetes)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
    for l in lignes:
        ws.append(list(l))
    for i, w in enumerate(largeurs):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A3"


def ecrire_modele(chemin, ref):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    p = ref["PARAMETRES"]
    _feuille(wb, "Paramètres", ["Paramètre", "Valeur"],
             [("Visites par intervenant et par jour", p["visites_par_intervenant_jour"]),
              ("Seuil de visites pour compter un cabinet", p["seuil_visites_cabinet"]),
              ("Cabinets physiques CSMKL2", p["cabinets_physiques_csmkl2"])], [45, 12],
             "Modifier seulement la colonne Valeur.")
    _feuille(wb, "Lits CHME", ["Unité", "Lits"], ref["LITS"], [25, 10],
             "Une ligne par unité d'hospitalisation (libellés Hospi CHIR, Hospi GO...).")
    _feuille(wb, "MAISON ROSE", ["UF rattachée à MAISON ROSE (CSMKL2)"], [(u,) for u in ref["MAISON_ROSE_UF"]], [45],
             "Les autres UF de CSMKL2 restent en activité principale.")
    _feuille(wb, "UF visites", ["UF (export visites)", "Spécialité affichée"], sorted(ref["UF_SPECIALITE"].items()), [45, 35],
             "Ajouter ici une nouvelle UF signalée « à valider » dans Notez bien.")
    _feuille(wb, "Catégories Evo", ["Categorie_acte (Evolucare)", "Spécialité", "Sous-spécialité"],
             [(k, a, b) for k, (a, b) in sorted(ref["EVO_CATEGORIES"].items())], [45, 25, 30],
             "Harmonise les catégories d'actes Evolucare avec les spécialités GPS.")
    _feuille(wb, "Produits Evo", ["Categorie_acte traitée comme produit"], [(x,) for x in ref["EVO_PRODUITS"]], [40],
             "Lignes séparées des actes / prestations (médicaments, consommables).")
    _feuille(wb, "Alias médecins", ["Nom nettoyé source", "Nom regroupé"], sorted(ref["ALIAS_MEDECINS"].items()), [35, 35],
             "Alias confirmés. Écrire les noms en MAJUSCULES sans accents ni titre (Dr, Pr...).")
    wb.save(chemin)


def _lignes(ws):
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row and row[0] not in (None, ""):
            yield row


def charger(chemin):
    """Renvoie le référentiel (défauts + fichier Excel). Crée le fichier s'il n'existe pas."""
    ref = defauts()
    if not os.path.exists(chemin):
        try:
            ecrire_modele(chemin, ref)
        except OSError:
            pass
        return ref
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    noms = wb.sheetnames
    if "Paramètres" in noms:
        cles = ["visites_par_intervenant_jour", "seuil_visites_cabinet", "cabinets_physiques_csmkl2"]
        for k, row in zip(cles, _lignes(wb["Paramètres"])):
            if isinstance(row[1], (int, float)):
                ref["PARAMETRES"][k] = int(row[1])
    if "Lits CHME" in noms:
        lits = [(str(r[0]).strip(), int(r[1])) for r in _lignes(wb["Lits CHME"]) if isinstance(r[1], (int, float))]
        if lits:
            ref["LITS"] = lits
    if "MAISON ROSE" in noms:
        ref["MAISON_ROSE_UF"] = [str(r[0]).strip() for r in _lignes(wb["MAISON ROSE"])]
    if "UF visites" in noms:
        ref["UF_SPECIALITE"] = {str(r[0]).strip(): str(r[1]).strip() for r in _lignes(wb["UF visites"]) if r[1]}
    if "Catégories Evo" in noms:
        ref["EVO_CATEGORIES"] = {str(r[0]).strip(): (str(r[1]).strip(), str(r[2]).strip())
                                 for r in _lignes(wb["Catégories Evo"]) if r[1] and r[2]}
    if "Produits Evo" in noms:
        ref["EVO_PRODUITS"] = [str(r[0]).strip() for r in _lignes(wb["Produits Evo"])]
    if "Alias médecins" in noms:
        ref["ALIAS_MEDECINS"] = {str(r[0]).strip().upper(): str(r[1]).strip().upper()
                                 for r in _lignes(wb["Alias médecins"]) if r[1]}
    wb.close()
    return ref
