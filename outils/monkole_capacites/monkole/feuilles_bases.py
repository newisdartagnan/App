"""Feuilles masquées : paramètres, bases de traçabilité et séries des graphiques."""
from openpyxl.utils import get_column_letter as L

from . import visites as V
from .styles import BLANC, CLAIR, DATE, DEC, NAVY, ecrire, mise_en_page
from .texte import nom_jour

ONGLET = "657A88"


def _feuille(wb, nom, entetes, largeurs=None):
    ws = wb.create_sheet(nom)
    mise_en_page(ws, ONGLET, figer="A6")
    ws.freeze_panes = "A5"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(6, len(entetes)))
    ecrire(ws, "A1", nom, taille=14, gras=True, couleur=BLANC, fond=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(6, len(entetes)))
    ecrire(ws, "A2", "Base de calcul / traçabilité — ne pas coller de nouvelles données sans réactualisation", fond=CLAIR)
    for i, t in enumerate(entetes, start=1):
        ecrire(ws, (5, i), t, gras=True, couleur=BLANC, fond=NAVY)
        ws.column_dimensions[L(i)].width = (largeurs[i - 1] if largeurs else 16)
    ws.sheet_state = "hidden"
    return ws


def _lignes(ws, lignes, formats=None, debut=6):
    formats = formats or {}
    for r, vals in enumerate(lignes, start=debut):
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in formats:
                cell.number_format = formats[c]
    if lignes:
        ws.auto_filter.ref = f"A5:{L(len(lignes[0]))}{debut + len(lignes) - 1}"


def parametres(wb, R):
    ws = wb.create_sheet("_Paramètres")
    mise_en_page(ws, ONGLET, figer="A5")
    ecrire(ws, "A1", "PARAMÈTRES REPRIS DU MODÈLE", taille=14, gras=True, couleur=BLANC, fond=NAVY)
    ecrire(ws, "A2", "Hypothèses, pas normes cliniques ni inventaire validé à cette date", fond=CLAIR)
    lignes = [("Début", R.debut), ("Fin", R.fin), ("Jours", R.ndays), ("Visites / intervenant / jour", R.cap_jour),
              ("Seuil de visites pour compter un cabinet", R.seuil), ("Cabinets physiques CSMKL2", R.cabinets_csmkl2)]
    r = 5
    for k, v in lignes:
        ws.cell(row=r, column=1, value=k)
        c = ws.cell(row=r, column=2, value=v)
        if hasattr(v, "year"):
            c.number_format = DATE
        r += 1
    r += 2
    for u, n in R.lits:
        ws.cell(row=r, column=1, value=u)
        ws.cell(row=r, column=2, value=n)
        r += 1
    ws.cell(row=r, column=1, value="Total lits CHME")
    ws.cell(row=r, column=2, value=R.total_lits)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    ws.sheet_state = "hidden"


def base_visites(wb, R):
    ws = _feuille(wb, "Base visites", ["Logiciel", "Ligne source", "Date / heure", "Jour", "Site", "Num_Dossier", "Clé dossier",
                                        "Intervenant source", "Intervenant regroupé", "Nature source", "Motif source", "UF source",
                                        "Activité", "Spécialité / unité", "Classe", "Volume retenu", "Champ de classement"])
    lignes = [(v["logiciel"], v["ligne"], v["date"], v["jour"], v["site"], v["dossier"], v["cle"], v["medecin_source"],
               v["medecin"], v["nature"], v["motif"], v["uf"], v["activite"], v["specialite"], v["classe"], 1, v["champ"])
              for v in R.visites]
    _lignes(ws, lignes, {3: "dd/mm/yyyy hh:mm", 4: DATE})


def base_hospitalisations(wb, R):
    ws = _feuille(wb, "Base hospitalisations", ["Logiciel", "Ligne source", "Site retenu", "Num_Dossier", "Entrée source",
                                                 "Sortie source", "UF / UH source", "Unité harmonisée", "Origine du site",
                                                 "Intervalle dans période", "Sortie après la fin", "Dates invalides"])
    lignes = [(b["logiciel"], b["ligne"], b["site"], b["dossier"], b["entree"], b["sortie"], b["uf"], b["unite"],
               b["origine_site"], b["dans_periode"], b["sortie_apres"], b["invalide"]) for b in R.base_hospi]
    _lignes(ws, lignes, {5: "dd/mm/yyyy hh:mm", 6: "dd/mm/yyyy hh:mm"})


def dossiers_hospitaliers(wb, R):
    jours = [j.strftime("%d/%m") for j in R.jours]
    ws = _feuille(wb, "Dossiers hospitaliers", ["Clé dossier", "Logiciel", "Site", "Num_Dossier", "Unité retenue", "Origine unité",
                                                 "Entrée source", "Sortie source", "Date entrée connue", "Jours connus",
                                                 "Entrée période", "Sortie période", "En cours à la fin", "Traces visites",
                                                 "Lignes Hospi", "Intervalles consolidés", "UF contradictoires"] + jours)
    lignes = []
    for d in R.dossiers:
        lignes.append([d["cle"], d["logiciel"], d["site"], d["dossier"], d["unite"], d["origine_unite"], d["entree"], d["sortie"],
                       d["date_connue"], d["jours_connus"], d["entree_periode"], d["sortie_periode"], d["en_cours"],
                       d["n_traces"], d["n_lignes"], len(d["intervalles"]), d["uf_contradictoires"]]
                      + ([round(x, 6) for x in d["par_jour"]] if d["date_connue"] else [None] * R.ndays))
    fm = {7: "dd/mm/yyyy hh:mm", 8: "dd/mm/yyyy hh:mm", 10: DEC}
    fm.update({18 + i: DEC for i in range(R.ndays)})
    _lignes(ws, lignes, fm)


def base_actes(wb, R):
    ws = _feuille(wb, "Base actes", ["Logiciel", "Ligne source", "Site", "Num_Dossier", "Clé dossier", "Date_V GPS",
                                      "DATEHEURE Evolucare", "Date visite retrouvée (contrôle)", "Date de la série source",
                                      "Catégorie source", "Spécialité harmonisée", "Sous-spécialité harmonisée", "Libellé acte",
                                      "Code acte Evo", "Quantité Evo", "Produit séparé", "Ligne retenue", "Profil dossier",
                                      "Consultation / résultat", "Répétition identique en plus", "Situation", "Dans la période"])
    lignes = [(a["logiciel"], a["ligne"], a["site"], a["dossier"], a["cle"], a["date_v"], a["dateheure"], a["date_visite"],
               a["date"], a["categorie"], a["specialite"], a["sous_specialite"], a["libelle"], a["code"], a["quantite"],
               a["produit"], a["retenue"], a["profil"], a["consultation"], a["repetition"], a["situation"], a["dans_periode"])
              for a in R.actes_tous]
    fmt = "dd/mm/yyyy hh:mm"
    _lignes(ws, lignes, {6: fmt, 7: fmt, 8: fmt, 9: fmt})


def medecins_jour(wb, R):
    ws = _feuille(wb, "_Médecins jour", ["Site", "Activité", "Jour", "Intervenant", "Consultation", "Résultats avant une semaine",
                                          "Résultats après une semaine", "Total visites", "GPS", "Evolucare", "Cabinet compté",
                                          "Capacité", "Marge cumulable", f"Au-delà de {R.cap_jour}", "Visite isolée"])
    lignes = []
    for (site, act), md in R.medjour.items():
        for (j, m) in sorted(md, key=lambda k: (k[0], k[1])):
            c = md[(j, m)]
            lignes.append((site, act, j, m, c["cons"], c["avant"], c["apres"], c["total"], c["gps"], c["evo"], c["cab"],
                           c["cap"], c["marge"], c["audela"], c["isolee"]))
    _lignes(ws, lignes, {3: DATE})


def jours(wb, R):
    """Séries quotidiennes (sources des graphiques). Ordre : CSMKL2 principale, MAISON ROSE, CHME ambulatoire, urgences."""
    ws = _feuille(wb, "_Jours", ["Site", "Activité", "Jour", "Jour affiché", "Visites", "GPS", "Evolucare", "Consultation",
                                  "Avant une semaine", "Après une semaine", "Cabinets comptés", "Capacité", "Utilisation",
                                  "Marge individuelle", "Dépassements individuels", "Sans médecin", "Visites isolées",
                                  "Jour au-delà capacité", "Cabinets toutes activités", "Étiquette"])
    lignes = []
    R.lignes_jours = {}
    r = 6
    for site, act in (("CSMKL2", V.PRINCIPALE), ("CSMKL2", V.SECONDAIRE), ("CHME", V.AMBULATOIRE), ("CHME", V.URGENCES)):
        R.lignes_jours[(site, act)] = r
        for l in R.jours_act[(site, act)]:
            j = l["jour"]
            lignes.append((site, act, j, f"{nom_jour(j)} {j.strftime('%d/%m')}", l["total"], l["gps"], l["evo"], l["cons"],
                           l["avant"], l["apres"], l["cab"], l["cap"], l["util"] if l["util"] is not None else "N/D",
                           l["marge"], l["depass"], l["sans"], l["isolees"], l["jour_audela"], R.cab_site[site][j],
                           j.strftime("%d/%m")))
            r += 1
    _lignes(ws, lignes, {3: DATE, 13: "0.0%"})


def hospi_jour(wb, R):
    ws = _feuille(wb, "_Hospi jour", ["Jour", "GPS journées / 8 unités", "Evo journées / 8 unités", "Lits CHME", "Entrées GPS",
                                       "Entrées Evo", "Sorties GPS", "Sorties Evo", "Dossiers sans entrée avec trace", "Étiquette"])
    lignes = [(h["jour"], round(h["gps"], 6), h["evo"], h["lits"], h["ent_gps"], h["ent_evo"], h["sor_gps"], h["sor_evo"],
               h["sans_trace"], h["jour"].strftime("%d/%m")) for h in R.hospi_jour]
    _lignes(ws, lignes, {1: DATE, 2: DEC, 3: DEC})
