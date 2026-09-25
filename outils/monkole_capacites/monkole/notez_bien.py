"""Feuille « Notez bien » : règles, sources, contrôles quantitatifs, rapprochements, lexique."""
import os
import re
from collections import defaultdict

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from . import visites as V
from .feuilles_detail import libelle_periode
from .feuilles_synthese import entete_tableau, fusion, section, zebre
from .styles import ALERTE_F, ALERTE_T, BLANC, CLAIR, DATE, GRIS, NB, ORANGE, ZEBRE, bandeau, ecrire, lien, mise_en_page
from .texte import MOIS, fr_nombre_espace as fr

DER = 15


def _nom_fichier(chemin):
    return re.sub(r"^[0-9a-f]{8}-", "", os.path.basename(chemin), flags=re.I)


def _texte(ws, r, titre, contenu):
    section(ws, r, titre, 1, DER)
    fusion(ws, r + 1, 1, DER)
    for c in range(1, DER + 1):
        ecrire(ws, (r + 1, c), contenu if c == 1 else None, couleur=GRIS, fond=CLAIR)
    ws.row_dimensions[r + 1].height = 47


def feuille_notez_bien(wb, R, positions, onglet):
    ws = wb.create_sheet("Notez bien")
    mise_en_page(ws, onglet, zoom=85, figer="A5")
    ws.column_dimensions["A"].width = 46
    for i in range(2, DER + 1):
        ws.column_dimensions[chr(64 + i)].width = 12
    per = f"{libelle_periode(R, '-')} {R.fin.year}"
    bandeau(ws, 1, "MONKOLE / RÈGLES, SOURCES ET CONTRÔLES", DER, taille=19, hauteur=32)
    fusion(ws, 2, 1, DER)
    ecrire(ws, "A2", f"{per} | Activités et capacités uniquement | Les diagnostics sont un travail séparé", fond=CLAIR)
    ws.row_dimensions[2].height = 28
    for c1, c2, t, cible in ((1, 5, "Dashboard", "'Dashboard'!A1"), (6, 10, "CSMKL2", "'CSMKL2'!A1"), (11, 15, "CHME", "'CHME'!A1")):
        fusion(ws, 3, c1, c2)
        lien(ws, (3, c1), t, cible)
    ws.row_dimensions[3].height = 24

    cap, seuil, cabs = R.cap_jour, R.seuil, R.cabinets_csmkl2
    d, f, fx = R.debut, R.fin, R.fin_excl
    fin_j = f"{f.day:02d}"
    t = R.unites_total
    sans_ent = t["sans_gps"] + t["sans_evo"]
    urg_sans = sum(l["sans"] for l in R.jours_act[("CHME", V.URGENCES)])
    urg_tot = sum(l["total"] for l in R.jours_act[("CHME", V.URGENCES)])
    evo_b = [b for b in R.base_hospi if b["logiciel"] == "Evolucare"]
    evo_dates_chme = sum(1 for b in evo_b if b["dans_periode"] and b["site"] == "CHME")
    evo_sans_site = sum(1 for b in evo_b if b["dans_periode"] and b["site"] not in ("CHME", "CSMKL2"))
    evo_hors = [b for b in evo_b if not b["dans_periode"]]
    dos_chme_evo = sum(1 for x in R.dossiers if x["logiciel"] == "Evolucare" and x["site"] == "CHME")
    dos_chme_gps = sum(1 for x in R.dossiers if x["logiciel"] == "GPS" and x["site"] == "CHME")
    sorties_evo = [b for b in evo_b if b["sortie"] is not None and b["dans_periode"]]
    toutes_apres = sorties_evo and all(b["sortie"] >= fx for b in sorties_evo)
    alias = "; ".join(f"{a} / {b}" for a, b in R.ref["ALIAS_MEDECINS"].items())
    vides = f"{f.day + 1:02d}-{_fin_mois(R):02d}" if f.day < _fin_mois(R) else "hors période"
    hors_txt = ""
    if evo_hors:
        mois_h = sorted({MOIS[b["entree"].month - 1] for b in evo_hors if b["entree"]})
        hors_txt = (f"Le dossier à entrée de {mois_h[0]} est conservé hors période. " if len(evo_hors) == 1 and mois_h
                    else f"{len(evo_hors)} dossier(s) à entrée hors période sont conservés hors période. ")
    lits_txt = ", ".join(str(n) for _, n in R.lits)

    blocs = [
        ("1 / Périmètre et actualisation",
         f"Les six nouveaux fichiers remplacent les anciens exports cumulés. Période : du {d:%d/%m/%Y} 00:00 au {fx:%d/%m/%Y} 00:00 exclu. "
         f"Deux sites documentés : CSMKL2 et CHME. Ne pas additionner d’anciens exports à ces fichiers {libelle_periode(R, '-')}."),
        ("2 / Structure conservée",
         "Un grand Dashboard commun, puis toutes les feuilles CSMKL2 ensemble, puis toutes les feuilles CHME ensemble. Médecins par jour "
         "et Jours par médecin présents pour chaque site; actes détaillés pour chacun. Les bases de contrôle sont masquées, pas supprimées."),
        ("3 / Une visite",
         "Une ligne conservée dans les fichiers Visites compte une visite enregistrée. Ce n’est pas un patient unique. Les traces "
         "hospitalières sont présentées à part, pas ajoutées aux visites ambulatoires ou d’urgence. Aucun acte n’ajoute une visite."),
        ("4 / Trois classes seulement",
         "Consultation; Résultats avant une semaine; Résultats après une semaine. Motif prioritaire; Nature utilisée quand Motif est "
         "vide/NULL. Tous les autres motifs restent dans Consultation conformément au modèle, y compris soins/pharmacie/examens "
         "enregistrés dans ce périmètre."),
        ("5 / CSMKL2 : activité principale et MAISON ROSE",
         "MAISON ROSE (ancien libellé : activités secondaires) : Consultation Prénatale Général, Imagerie/Exploration Imagerie, "
         "Laboratoire, Sage-Femme. Les autres UF restent principales. Aucune exclusion d’un médecin ayant une seule classe de motif : "
         "le modèle validé conserve toutes les visites."),
        ("6 / Cabinet compté, pas salle prouvée",
         f"Un intervenant nommé atteint au moins {seuil} visites dans une activité et une journée : 1 cabinet compté. GPS et Evolucare "
         "sont regroupés par nom normalisé. Les journées à une visite et les identités absentes restent dans les volumes mais ne créent "
         f"pas de cabinet. Les {cabs} cabinets CSMKL2 ne plafonnent pas le comptage; un dépassement appelle un contrôle de rotation."),
        ("7 / Journées de cabinet et capacité",
         "Une journée de cabinet = un cabinet compté pendant une journée. Exemple de définition : 3 cabinets pendant 2 jours = 6 "
         f"journées de cabinet. Capacité de période = {cap} × somme des cabinets comptés chaque jour. On ne multiplie pas une deuxième "
         f"fois par {R.ndays}."),
        ("8 / Utilisation et marge",
         "Utilisation du modèle = toutes les visites du périmètre / capacité comptée. Visites isolées et sans médecin restent au "
         f"numérateur. Marge individuelle cumulée = somme de MAX({cap} - visites, 0), uniquement pour les cabinets comptés. Cette "
         "marge n’est pas une garantie de places réellement disponibles."),
        ("9 / Identités des intervenants",
         f"Titres, accents et espaces sont harmonisés. Alias confirmés repris : {alias}. Les rapprochements uniques à au moins deux "
         "mots communs suivent la règle du modèle et sont listés plus bas pour contrôle. Aucun rapprochement sur un seul prénom/nom."),
        ("10 / Pic, minimum et maximum",
         "Les min/max individuels portent sur les jours avec au moins une visite dans le périmètre concerné. En cas de pic ex aequo, la "
         f"première date est retenue; les contributions GPS/Evo sont celles de cette même date. Les jours {vides} sont laissés vides "
         "dans les matrices, pas remplis de zéros."),
        ("11 / Urgences : manque d’identité",
         f"{urg_sans} visites d’urgences CHME n’ont pas de médecin renseigné. Elles restent comptées dans les {urg_tot} visites. Leur "
         "capacité individuelle ne peut pas être calculée. Le quotient visites/capacité des seuls intervenants nommés ne permet pas une "
         "décision fiable sur la capacité réelle des urgences."),
        ("12 / Spécialités et capacité physique",
         "Au CHME, la spécialité vient de l’UF, sinon du motif explicite. Hospitalisation est prioritaire lorsque UF/Motif/Nature la "
         "mentionnent; Urgences vient ensuite. Un intervenant couvrant plusieurs spécialités ne multiplie pas les cabinets de la même "
         "activité. Les repères par spécialité sont non additionnables."),
        ("13 / Actes : dates GPS et Evolucare",
         "GPS : Date_V est la seule date de filtrage et des courbes. Evolucare : aucune Date_V dans l’export; DATEHEURE est affichée dans "
         "une série complémentaire séparée. La date des visites retrouvée par dossier est conservée pour audit, jamais substituée "
         "silencieusement à Date_V. Aucun total de courbe GPS + Evo n’est présenté."),
        ("14 / Une ligne d’acte, pas une quantité universelle",
         "Les chiffres principaux des actes sont des enregistrements (une ligne chacun). Les quantités Evolucare sont conservées dans "
         "une colonne séparée; GPS n’a pas de quantité. Le nombre de séances cité dans un forfait ne multiplie pas la ligne GPS. Sans "
         "identifiant d’acte, les répétitions identiques sont signalées mais conservées."),
        ("15 / Produits et actes non classés",
         f"{fr(len(R.produits))} lignes Evolucare de médicament/consommable médical sont hors des actes/prestations; conservées dans "
         "Base actes. Les catégories vides ou ACTE restent sous Catégorie à préciser, sans attribution clinique inventée. Les "
         "consultations, résultats, nursing, hébergement et prestations hôtelières restent identifiables, pas transformés en examens "
         "techniques. « Monkole 3 » est une spécialité source, pas un troisième site déduit."),
        ("16 / Rapprochement des dossiers",
         "Clé = logiciel + site + Num_Dossier, conservé comme texte. Visites et hospitalisations sont consolidées avant rattachement : "
         "une seule ligne par acte, pas de produit cartésien. Profil exclusif : Hospitalisation, sinon Ambulatoire, sinon Urgences, "
         "sinon Sans correspondance. Ce profil ne prouve pas que l’acte s’est déroulé pendant le séjour."),
        ("17 / Actes des dossiers vus en consultation",
         "Dans CHME, Actes liés désigne tous les actes des dossiers passés dans la spécialité, pas seulement ceux produits par cette "
         "spécialité. Un dossier pouvant voir plusieurs spécialistes, ces colonnes et les dossiers distincts par spécialité ne "
         "s’additionnent pas."),
        ("18 / Dossiers hospitaliers",
         "Union des dossiers avec intervalle Hospi couvrant la période et des dossiers explicitement hospitaliers dans Visites, à "
         f"l’intérieur de chaque logiciel/site. Un acte seul ne crée aucun séjour. GPS : {dos_chme_gps} dossiers CHME; Evo : "
         f"{dos_chme_evo} CHME et {len(R.sans_site)} sans site. {sans_ent} dossiers CHME sans date d’entrée restent comptés, sans "
         "durée ajoutée."),
        ("19 / Site des hospitalisations Evolucare",
         "Le fichier Hospi Evo ne fournit pas de site. Site attribué uniquement lorsqu’un Num_Dossier retrouve un site unique cohérent "
         f"dans Visites/Actes Evo : {evo_dates_chme} séjours datés CHME. {evo_sans_site} séjours couvrant {MOIS[d.month - 1]} restent "
         f"sans site. {hors_txt}Aucune unité ou personne n’est utilisée pour deviner un site."),
        ("20 / Durées et sorties",
         f"Intervalles [entrée, sortie[ limités à la période; répétitions/chevauchements neutralisés avant somme. Une sortie vide ou "
         f"postérieure au {fin_j} laisse le dossier en cours au {fin_j}. Une date après le {fin_j} n’est pas une sortie réalisée dans la "
         "période. Les dates Evo ne portent pas d’heures : début/fin à minuit selon l’export, précision à confirmer. "
         + (f"Toutes les sorties Evo fournies sont postérieures au {fin_j} : zéro sortie dans cet export ne prouve pas zéro sortie "
            "réelle. " if toutes_apres else "")
         + "Nature prévisionnelle des dates à confirmer."),
        ("21 / Occupation : deux sources séparées",
         f"GPS : journées connues des {'huit' if len(R.lits) == 8 else len(R.lits)} unités / ({R.total_lits} × {R.ndays}). Evo : même dénominateur mais série distincte, "
         "sur les dossiers datés dont le site est retrouvé. Les taux ne sont pas additionnés : pas d’identifiant patient commun "
         "permettant de dédoublonner les séjours interlogiciels. Les jours hors des " f"{'huit' if len(R.lits) == 8 else len(R.lits)} unités restent hors taux."),
        ("22 / Occupation incomplète",
         "Aucun jour de séjour n’est inventé pour un dossier sans entrée. Les actes ne servent pas à lui attribuer des jours. Les séjours "
         f"antérieurs au {d:%d/%m} sont intégrés seulement si leurs dates sont effectivement fournies. Le complément d’un taux à 100 % "
         "ne prouve pas que les lits sont libres."),
        ("23 / Unités et lits",
         "UF Hospi prioritaire, UH Evo rapprochée par libellé d’unité : chirurgie, gynéco-obstétrique, médecine interne, néonatologie, "
         "pédiatrie, urgences, soins intensifs pédiatriques, réanimation. Pas de réattribution à partir du diagnostic. Hors unité / UF "
         f"absente reste hors capacité. Inventaire repris : {lits_txt} lits."),
        ("24 / Limites et diffusion",
         "Les exports ne prouvent pas l’exhaustivité de l’hôpital. Ils ne donnent ni horaires/cabinets physiques CHME ni capacité des "
         "appareils. Pas d’incidence médicale, de taux de saturation technique ou de patients uniques inventés. Base sans adresses ni "
         "contacts; Num_Dossier et noms d’intervenants réservés aux personnes habilitées."),
        ("25 / Calculs et prochaine mise à jour",
         f"Les calculs et leurs résultats sont enregistrés. Ce classeur est un arrêté contrôlé au {f:%d/%m}; une prochaine extraction "
         "se traite en relançant le programme sur les six nouveaux exports, qui reconstruit les liens et dédoublonnages. Le fichier "
         "diagnostics et son dictionnaire ne sont pas modifiés par ce travail."),
    ]
    r = 5
    for titre, contenu in blocs:
        _texte(ws, r, titre, contenu)
        r += 3

    # Contrôles quantitatifs
    r = 80
    section(ws, r, "CONTRÔLES QUANTITATIFS / LES ÉCARTS ATTENDUS SONT ZÉRO", 1, DER)
    entete_tableau(ws, r + 1, [(1, 1, "Contrôle"), (2, 2, "Calcul"), (3, 3, "Attendu"), (4, 4, "Écart")])
    for c in range(5, DER + 1):
        ecrire(ws, (r + 1, c), fond=None)
    n_vg, n_ve = len(R.donnees["visites_gps"]), len(R.donnees["visites_evo"])
    hors_hosp = sum(1 for v in R.visites if v["activite"] != V.HOSPITALISATION)
    traces = sum(1 for v in R.visites if v["activite"] == V.HOSPITALISATION)
    cs_j = sum(l["total"] for a in (V.PRINCIPALE, V.SECONDAIRE) for l in R.jours_act[("CSMKL2", a)])
    ch_j = sum(l["total"] for a in (V.AMBULATOIRE, V.URGENCES) for l in R.jours_act[("CHME", a)])
    arbres = positions["arbres"]

    def arbre_tot(site, k):
        return sum(sp["agg"][k] for sp in arbres[site])

    base = lambda lg, site: sum(1 for a in R.actes if a["logiciel"] == lg and a["site"] == site)
    ctrl = [
        ("Visites source GPS + Evolucare", len(R.visites) + len(R.visites_hors_periode), n_vg + n_ve),
        ("Hors hospi + traces = toutes les visites", hors_hosp + traces, len(R.visites)),
        ("Visites CSMKL2", cs_j, sum(1 for v in R.visites if v["site"] == "CSMKL2")),
        ("Visites CHME hors hospitalisation", ch_j, sum(1 for v in R.visites if v["site"] == "CHME" and v["activite"] != V.HOSPITALISATION)),
        ("GPS : toutes les lignes actes conservées", sum(1 for a in R.actes_tous if a["logiciel"] == "GPS"), len(R.donnees["actes_gps"])),
        ("Evo : prestations + produits = export", sum(1 for a in R.actes_tous if a["logiciel"] == "Evolucare"), len(R.donnees["actes_evo"])),
        ("CSMKL2 actes GPS", arbre_tot("CSMKL2", "gps"), base("GPS", "CSMKL2")),
        ("CSMKL2 actes Evolucare", arbre_tot("CSMKL2", "evo"), base("Evolucare", "CSMKL2")),
        ("CHME actes GPS", arbre_tot("CHME", "gps"), base("GPS", "CHME")),
        ("CHME actes Evolucare", arbre_tot("CHME", "evo"), base("Evolucare", "CHME")),
        ("Hospi : CHME + sans site", len(R.dossiers), sum(1 for x in R.dossiers if x["site"] == "CHME") + len(R.sans_site)),
        ("GPS : journées connues unités + hors unité", sum(u["jours_gps"] for u in R.unites),
         sum(x["jours_connus"] or 0 for x in R.dossiers if x["logiciel"] == "GPS" and x["site"] == "CHME")),
        ("Jours de période", R.ndays, (R.fin - R.debut).days + 1),
    ]
    R.controles = ctrl
    for i, (lib, a, b) in enumerate(ctrl):
        rr = r + 2 + i
        ecrire(ws, (rr, 1), lib, fond=zebre(rr))
        for c, v in ((2, a), (3, b), (4, a - b)):
            ecrire(ws, (rr, c), v, fond=zebre(rr), fmt=NB if isinstance(v, int) else "#,##0.00", h="right")
        for c in range(5, DER + 1):
            ecrire(ws, (rr, c), fond=zebre(rr), h="right")
        ws.row_dimensions[rr].height = 28
    fin_ctrl = r + 1 + len(ctrl)
    ws.conditional_formatting.add(f"D{r + 2}:D{fin_ctrl}",
                                  CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor=ORANGE)))

    # Sources
    r = fin_ctrl + 4
    section(ws, r, "SOURCES UTILISÉES ET CONTRÔLE DES DATES", 1, DER)
    entete_tableau(ws, r + 1, [(1, 6, "Fichier"), (7, 7, "Lignes source"), (8, 8, "Début observé"), (9, 9, "Fin observée"),
                               (10, 15, "Précaution")], hauteur=38)
    n_dates_gps = sum(1 for x in R.dossiers if x["logiciel"] == "GPS" and x["date_connue"])

    def bornes(lignes, cle):
        ds = [x[cle] for x in lignes if x.get(cle) is not None]
        ds = [x if hasattr(x, "year") else None for x in ds]
        ds = [x for x in ds if x]
        if not ds:
            return None, None
        a, b = min(ds), max(ds)
        return a.replace(hour=0, minute=0, second=0, microsecond=0), b.replace(hour=0, minute=0, second=0, microsecond=0)

    D = R.donnees
    sources = [
        ("visites_gps", "Date", "Date des visites"),
        ("visites_evo", "Date", "Date des visites; motifs souvent vides"),
        ("actes_gps", "Date_V", "Date_V uniquement"),
        ("actes_evo", "DATEHEURE", "DATEHEURE, complément distinct; produits séparés"),
        ("hospi_gps", "Date Entrée", f"{n_dates_gps} dossiers datés GPS après consolidation"),
        ("hospi_evo", "Date_Entree", f"{sum(1 for b in evo_b if b['dans_periode'])} intervalles couvrent {MOIS[d.month - 1]}; "
                                     f"{len(evo_hors)} entrée hors période"),
    ]
    for i, (k, col, prec) in enumerate(sources):
        rr = r + 2 + i
        a, b = bornes(D[k], col)
        fusion(ws, rr, 1, 6)
        ecrire(ws, (rr, 1), _nom_fichier(R.chemins[k]), fond=BLANC if i % 2 == 0 else ZEBRE)
        ecrire(ws, (rr, 7), len(D[k]), fmt=NB, h="right", fond=BLANC if i % 2 == 0 else ZEBRE)
        for c, v in ((8, a), (9, b)):
            cc = ws.cell(row=rr, column=c, value=v)
            cc.number_format = DATE
        fusion(ws, rr, 10, 15)
        ecrire(ws, (rr, 10), prec, couleur=GRIS, fond=BLANC if i % 2 == 0 else ZEBRE)
        ws.row_dimensions[rr].height = 30
    rr = r + 2 + len(sources) + 2
    fusion(ws, rr, 1, DER)
    ecrire(ws, (rr, 1), "Modèle reproduit : Monkole_Activites_Capacites_01-15_Septembre_2026_Version_retenue.xlsx (onglets par site, "
                        "règles de capacité CSMKL2 et CHME, noms, couleurs et navigation). Classeur produit hors connexion par le "
                        "programme Monkole Capacités.", couleur=GRIS, fond=CLAIR)
    ws.row_dimensions[rr].height = 35

    # Profils des actes
    r = rr + 3
    section(ws, r, "ACTES / RATTACHEMENT AUX DOSSIERS PAR SITE ET LOGICIEL", 1, DER)
    entete_tableau(ws, r + 1, [(1, 1, "Site"), (2, 2, "Logiciel"), (3, 3, "Profil dossier"), (4, 4, "Enregistrements"),
                               (5, 5, "Dossiers concernés")])
    grp = defaultdict(list)
    for a in R.actes:
        grp[(a["site"], a["logiciel"], a["profil"])].append(a)
    rr = r + 2
    for k in sorted(grp):
        vals = list(k) + [len(grp[k]), len({a["cle"] for a in grp[k]})]
        for c, v in enumerate(vals, start=1):
            ecrire(ws, (rr, c), v, fond=zebre(rr), fmt=NB if c >= 4 else None, h="right" if c >= 4 else "left")
        ws.row_dimensions[rr].height = 24
        rr += 1

    # Hospitalisations Evolucare sans site
    r = rr + 3
    section(ws, r, "HOSPITALISATIONS EVOLUCARE SANS SITE / NON ATTRIBUÉES", 1, DER)
    entete_tableau(ws, r + 1, [(1, 1, "Num_Dossier"), (2, 2, "Unité source"), (3, 3, "Entrée"), (4, 4, "Sortie source"), (5, 15, "Statut")])
    rr = r + 2
    for b in sorted(R.hospi_evo_non_attribues, key=lambda b: b["dossier"] or ""):
        if b["dans_periode"]:
            statut = "Couvre la période; site à compléter"
        else:
            statut = f"Hors période : entrée après le {f:%d/%m}, aucune correction supposée" if (b["entree"] and b["entree"] >= fx) \
                else "Hors période, aucune correction supposée"
        ecrire(ws, (rr, 1), b["dossier"], fond=zebre(rr))
        ecrire(ws, (rr, 2), b["uf"], fond=zebre(rr))
        ecrire(ws, (rr, 3), b["entree"], fond=zebre(rr), fmt=DATE, h="right")
        ecrire(ws, (rr, 4), b["sortie"], fond=zebre(rr), fmt=DATE, h="right")
        fusion(ws, rr, 5, DER)
        ecrire(ws, (rr, 5), statut, couleur=ALERTE_T, fond=ALERTE_F)
        ws.row_dimensions[rr].height = 31
        rr += 1

    # Noms rapprochés
    r = rr + 3
    section(ws, r, "NOMS RAPPROCHÉS / CONTRÔLE DES IDENTITÉS", 1, DER)
    entete_tableau(ws, r + 1, [(1, 5, "Nom nettoyé source"), (6, 10, "Nom regroupé"), (11, 15, "Règle de rapprochement")], hauteur=38)
    rr = r + 2
    for a, b in R.rapprochements:
        fusion(ws, rr, 1, 5)
        fusion(ws, rr, 6, 10)
        fusion(ws, rr, 11, 15)
        ecrire(ws, (rr, 1), a, fond=zebre(rr + 1))
        ecrire(ws, (rr, 6), b, fond=zebre(rr + 1))
        ecrire(ws, (rr, 11), "Deux mots identiques; correspondance unique dans les exports; à confirmer en cas d’homonymie.",
               taille=9, couleur=GRIS, fond=zebre(rr + 1))
        ws.row_dimensions[rr].height = 31
        rr += 1

    # Nouveautés à valider (seulement si nécessaire)
    nouveautes = [("UF de visite sans correspondance (spécialité = libellé UF)", k, n) for k, n in sorted(R.uf_inconnues.items())]
    nouveautes += [("Catégorie d’acte Evolucare nouvelle (découpée sur le tiret)", k, n) for k, n in sorted(R.categories_inconnues.items())]
    from .classeur_brut import AVERTISSEMENTS
    nouveautes += [("Export lu malgré un défaut de fichier", avis, None) for avis in AVERTISSEMENTS]
    if R.visites_hors_periode:
        nouveautes.append(("Visites hors période exclues", "Date hors période ou absente", len(R.visites_hors_periode)))
    if R.actes_hors_periode:
        nouveautes.append(("Actes hors période exclus des séries", "Date_V / DATEHEURE hors période ou absente", len(R.actes_hors_periode)))
    if nouveautes:
        r = rr + 3
        section(ws, r, "À VALIDER / NOUVEAUX LIBELLÉS OU LIGNES HORS PÉRIODE", 1, DER)
        entete_tableau(ws, r + 1, [(1, 5, "Point"), (6, 12, "Libellé"), (13, 15, "Lignes")])
        rr = r + 2
        for a, b, n in nouveautes:
            fusion(ws, rr, 1, 5)
            fusion(ws, rr, 6, 12)
            fusion(ws, rr, 13, 15)
            ecrire(ws, (rr, 1), a, fond=zebre(rr))
            ecrire(ws, (rr, 6), b, couleur=ALERTE_T, fond=ALERTE_F)
            ecrire(ws, (rr, 13), n, fmt=NB, h="right", fond=zebre(rr))
            ws.row_dimensions[rr].height = 31
            rr += 1

    # Lexique
    r = rr + 3
    section(ws, r, "PETIT LEXIQUE POUR LIRE LE DASHBOARD", 1, DER)
    lex = [("N/D", "Non disponible ou non calculable; ce n’est pas zéro."),
           ("Jours actifs", "Jours avec au moins une visite enregistrée, pas nombre de jours de présence réelle au travail."),
           ("Enregistrements d’actes", "Lignes d’activité conservées; ne pas confondre avec doses, séances ou patients."),
           ("Dossiers concernés", "Num_Dossier distincts dans chaque logiciel et site; pas de patients uniques interlogiciels."),
           ("Profil du dossier", "Activité retrouvée par les liens : hospitalisation, ambulatoire ou urgences. Ne prouve pas la "
                                 "localisation de chaque acte."),
           ("Occupation documentée", "Part des journées-lits expliquée par les dates connues; pas occupation exhaustive ni mesure des "
                                     "lits libres.")]
    for i, (a, b) in enumerate(lex):
        rr = r + 1 + i
        fusion(ws, rr, 1, 4)
        fusion(ws, rr, 5, DER)
        ecrire(ws, (rr, 1), a, gras=True, fond=zebre(rr + 1))
        ecrire(ws, (rr, 5), b, couleur=GRIS, fond=zebre(rr + 1))
        ws.row_dimensions[rr].height = 31
    return ws


def _fin_mois(R):
    import calendar
    return calendar.monthrange(R.fin.year, R.fin.month)[1]
