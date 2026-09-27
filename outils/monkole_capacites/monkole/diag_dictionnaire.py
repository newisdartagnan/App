"""Dictionnaire des diagnostics : formulations exactes -> diagnostics regroupés -> familles.

Fichier modifiable : Dictionnaire_Diagnostics.xlsx (à côté du programme).
  - « Formulations » : une ligne par (ID, diagnostic regroupé). Une phrase composite occupe plusieurs lignes
    avec le même ID. « Versions exactes » contient les textes d'origine, un par ligne de cellule.
  - « Groupes » : code du diagnostic regroupé, libellé, famille, nature.
"""
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .texte import cle

STATUTS = ["Sans marqueur de doute", "Hypothèse explicite", "À clarifier", "Contexte / antécédent", "Non exploitable",
           "Explicitement écarté"]
DECRIT = "Sans marqueur de doute"
HYPOTHESE = "Hypothèse explicite"

FAMILLES = ["Grossesse / accouchement", "ORL / respiratoire", "Cardiovasculaire", "Endocrino-métabolique",
            "Infectieux / parasitaire", "Ophtalmologie", "Digestif / hépatique", "Urinaire / néphrologie",
            "Locomoteur / traumatologie", "Gynécologie / sein", "Neurologie / santé mentale", "Peau / allergie",
            "Hématologie / oncologie", "Néonatal / périnatal", "Signes généraux / iatrogénie", "Bucco-dentaire"]
HORS_FAMILLES = ["À clarifier", "Contexte / hors diagnostic"]

COLS_FORM = ["ID", "Libellé représentatif", "Code groupe", "Statut dans le texte", "Méthode / prudence",
             "Validation médicale", "Commentaire de validation", "Versions exactes"]
COLS_GROUPES = ["Code groupe", "Diagnostic regroupé", "Famille", "Nature du groupe"]

# Groupe réservé aux formulations nouvelles non encore classées
CODE_NOUVEAU = "QNEW"
GROUPE_NOUVEAU = ("Formulation nouvelle à classer", "À clarifier", "Qualité / ambiguïté")

ENTETE = PatternFill("solid", fgColor="143348")


class Dictionnaire:
    def __init__(self):
        self.groupes = {}          # code -> (libellé, famille, nature)
        self.entrees = {}          # ID -> {"libelle", "lignes": [(code, statut, methode, validation, commentaire)], "versions": [..]}
        self.exact = {}            # texte exact -> ID
        self.normal = {}           # clé normalisée -> ID

    def indexer(self):
        self.exact, self.normal = {}, {}
        for i, e in self.entrees.items():
            for v in e["versions"]:
                self.exact.setdefault(v, i)
                self.normal.setdefault(cle_texte(v), i)

    def prochain_id(self):
        n = max((int(m.group(1)) for i in self.entrees for m in [re.fullmatch(r"D(\d+)", i)] if m), default=0)
        return f"D{n + 1:04d}"


def cle_texte(t):
    """Normalisation prudente : casse, accents, espaces et ponctuation d'extrémité."""
    if t is None:
        return ""
    k = cle(str(t))
    return k.strip(" .,;:-")


def charger(chemin):
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    d = Dictionnaire()
    for row in list(wb["Groupes"].iter_rows(min_row=3, values_only=True)):
        if row and row[0]:
            d.groupes[str(row[0]).strip()] = (str(row[1]).strip(), str(row[2]).strip(), str(row[3]).strip())
    for row in wb["Formulations"].iter_rows(min_row=3, values_only=True):
        if not row or not row[0] or not row[2]:
            continue
        i = str(row[0]).strip()
        e = d.entrees.setdefault(i, {"libelle": row[1], "lignes": [], "versions": []})
        e["lignes"].append((str(row[2]).strip(), str(row[3] or DECRIT).strip(), row[4], row[5], row[6]))
        if row[7] is not None:
            for v in str(row[7]).split("\n"):
                if v not in e["versions"]:
                    e["versions"].append(v)
        elif "" not in e["versions"] and row[1] in (None, ""):
            e["versions"].append("")          # entrée des diagnostics vides
    wb.close()
    d.groupes.setdefault(CODE_NOUVEAU, GROUPE_NOUVEAU)
    d.indexer()
    return d


def _feuille(wb, nom, aide, entetes, lignes, largeurs):
    ws = wb.create_sheet(nom)
    ws.append([aide])
    ws["A1"].font = Font(italic=True, color="657A88")
    ws.append(entetes)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for l in lignes:
        ws.append(list(l))
    for i, w in enumerate(largeurs):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{chr(64 + len(entetes))}{ws.max_row}"
    return ws


def ecrire(chemin, d):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    lignes = []
    for i in sorted(d.entrees):
        e = d.entrees[i]
        versions = "\n".join(e["versions"]) if e["versions"] != [""] else None
        for code, statut, methode, validation, commentaire in e["lignes"]:
            lignes.append((i, e["libelle"], code, statut, methode, validation, commentaire, versions))
    ws = _feuille(wb, "Formulations",
                  "Une ligne par formulation et par diagnostic regroupé. Phrase composite = plusieurs lignes avec le même ID. "
                  "Versions exactes : les textes d'origine, un par ligne (Alt+Entrée).",
                  COLS_FORM, lignes, [9, 40, 12, 22, 45, 14, 30, 60])
    for row in ws.iter_rows(min_row=3):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    n = ws.max_row
    dv = DataValidation(type="list", formula1="=Groupes!$A$3:$A$2000", allow_blank=False)
    dv.add(f"C3:C{n + 2000}")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=False)
    dv2.add(f"D3:D{n + 2000}")
    ws.add_data_validation(dv2)
    _feuille(wb, "Groupes", "Diagnostics regroupés. La famille détermine la place dans le disque (familles À clarifier et "
                            "Contexte / hors diagnostic : hors disque).",
             COLS_GROUPES, [(c,) + d.groupes[c] for c in sorted(d.groupes)], [14, 55, 30, 22])
    aide = wb.create_sheet("Aide")
    for t in ["Ajouter une nouvelle formulation :",
              "1. Ouvrir le fichier sorties/Diagnostics_a_classer_… produit par le programme.",
              "2. Vérifier ou corriger la colonne Code groupe (liste de la feuille Groupes) et le Statut.",
              "3. Copier les lignes validées à la fin de la feuille Formulations, puis relancer le programme.",
              "Créer un nouveau diagnostic regroupé : ajouter une ligne dans Groupes (code court en majuscules).",
              "Statuts : " + " ; ".join(STATUTS) + ".",
              "Seul « Sans marqueur de doute » entre dans les diagnostics comptabilisés ; « Hypothèse explicite » "
              "alimente les hypothèses seules ; les autres statuts restent hors du disque."]:
        aide.append([t])
    aide.column_dimensions["A"].width = 120
    wb.save(chemin)
